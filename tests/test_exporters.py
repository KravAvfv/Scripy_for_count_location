from __future__ import annotations

import os
import zipfile
from pathlib import Path

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook

from sitesizer.core.catalog import Catalog
from sitesizer.core.presets import preset
from sitesizer.core.sizing import size_site
from sitesizer.exporters.diagram import ProjectDiagram, SiteDiagram, export_png, export_svg
from sitesizer.exporters.pdf import export_pdf
from sitesizer.exporters.xlsx import export_xlsx

from .conftest import make_site


@pytest.fixture(scope="module")
def hq_result(catalog: Catalog):
    p = preset("hq")
    assert p is not None
    return size_site(p.build(name="Тест HQ"), catalog)


def test_xlsx_structure(tmp_path: Path, catalog: Catalog, hq_result) -> None:
    path = export_xlsx(hq_result, tmp_path / "out.xlsx", catalog=catalog, options={"author": "A", "project": "P"})
    wb = load_workbook(path)
    name = hq_result.input.name
    assert wb.sheetnames == [
        "Слаботрумка",
        f"{name} - Схема+шафи"[:31],
        f"{name} - IP",
        "Ціни та кількість",
        "Вихідні дані",
    ]
    with zipfile.ZipFile(path) as z:
        assert any(n.startswith("xl/media/") for n in z.namelist())


def test_xlsx_spec_matches_template(tmp_path: Path, catalog: Catalog) -> None:
    r = size_site(make_site(sockets=96, cameras=10), catalog)
    wb = load_workbook(export_xlsx(r, tmp_path / "s.xlsx", catalog=catalog, options={"sheets": ["spec"]}))
    assert wb.sheetnames == ["Слаботрумка"]
    ws = wb["Слаботрумка"]
    assert [ws.cell(row=4, column=c).value for c in range(1, 9)] == [
        "Код", "Найменування", "Од.", "К-сть", "Ціна (E)", "Сума (E)", "Ціна", "Сума"
    ]  # fmt: skip
    sections = [
        ws.cell(row=r, column=1).value
        for r in range(5, ws.max_row + 1)
        if ws.cell(row=r, column=1).fill.start_color.rgb.endswith("BFBFBF")
    ]
    assert sections == ["СКС", "Network"]
    rows = {ws.cell(row=r, column=2).value: r for r in range(5, ws.max_row + 1)}
    row = next(r for name, r in rows.items() if name and name.startswith("FS-148F —"))
    assert ws.cell(row=row, column=1).value == "000084545" and ws.cell(row=row, column=4).value == 2
    assert (
        ws.cell(row=row, column=6).value == f"=D{row}*E{row}" and ws.cell(row=row, column=8).value == f"=D{row}*G{row}"
    )
    unused = next(r for name, r in rows.items() if name and name.startswith("FS-648F —"))
    assert ws.cell(row=unused, column=4).value == 0  # every template item is listed, 0 when unused
    only = load_workbook(
        export_xlsx(r, tmp_path / "o.xlsx", catalog=catalog, options={"sheets": ["spec"], "only_used": True})
    )["Слаботрумка"]
    assert all(
        (only.cell(row=i, column=4).value or 0) > 0
        for i in range(5, only.max_row + 1)
        if only.cell(row=i, column=6).value and str(only.cell(row=i, column=6).value).startswith("=D")
    )


def test_xlsx_prices(tmp_path: Path, catalog: Catalog) -> None:
    r = size_site(make_site(sockets=10), catalog)
    for line in r.bom:
        if line.qty:
            line.unit_price = 100.0
    wb = load_workbook(export_xlsx(r, tmp_path / "p.xlsx", catalog=catalog, options={"vat_pct": 20}))
    ws = wb["Ціни та кількість"]
    assert ws.cell(row=4, column=8).value == "Сума"
    assert all("Обґрунтування" not in str(c.value) for row in ws.iter_rows() for c in row)


def test_pdf(tmp_path: Path, catalog: Catalog, hq_result) -> None:
    path = export_pdf(hq_result, tmp_path / "r.pdf", catalog=catalog, options={"customer": "X"})
    data = path.read_bytes()
    assert data.startswith(b"%PDF") and len(data) > 20_000


def test_diagram_exports(tmp_path: Path, catalog: Catalog, hq_result) -> None:
    png = export_png(SiteDiagram(hq_result, catalog), tmp_path / "d.png")
    assert png.read_bytes()[:4] == b"\x89PNG"
    svg = export_svg(ProjectDiagram([("HQ", True, hq_result), ("B", False, hq_result)], catalog), tmp_path / "d.svg")
    assert "<svg" in svg.read_text(encoding="utf-8")


def test_rack_exports(tmp_path: Path, catalog: Catalog, hq_result) -> None:
    from sitesizer.exporters.rack import RackDiagram, rack_table_rows

    d = RackDiagram(hq_result, catalog)
    assert d.size().height() > hq_result.rack.plans[0].size_u * d.U_H
    assert export_png(d, tmp_path / "r.png").read_bytes()[:4] == b"\x89PNG"  # type: ignore[arg-type]
    assert "<svg" in export_svg(d, tmp_path / "r.svg").read_text(encoding="utf-8")  # type: ignore[arg-type]
    rows = rack_table_rows(hq_result)
    assert len(rows) == sum(len(p.items) for p in hq_result.rack.plans)
    wb = load_workbook(export_xlsx(hq_result, tmp_path / "x.xlsx", catalog=catalog, options={"sheets": ["racks"]}))
    ws = wb.worksheets[0]
    plan = hq_result.rack.plans[0]
    assert str(ws.cell(row=3, column=2).value).startswith(plan.name)
    assert ws.cell(row=5, column=2).value == plan.size_u and ws.cell(row=5, column=4).value == plan.size_u
    labels = [ws.cell(row=r, column=3).value for r in range(5, 5 + plan.size_u)]
    assert "Органайзер" in labels


def test_empty_diagram(catalog: Catalog) -> None:
    d = SiteDiagram(size_site(make_site(), catalog), catalog)
    assert d.empty and d.size().width() > 0


def test_cli_exports(tmp_path: Path) -> None:
    from sitesizer import cli

    inp = tmp_path / "s.json"
    inp.write_text('{"name": "X", "sockets": 30, "cameras": 4, "mode": "extended"}', encoding="utf-8")
    assert (
        cli.main(["--json", str(inp), "--xlsx", str(tmp_path / "x.xlsx"), "--pdf", str(tmp_path / "x.pdf"), "--quiet"])
        == 0
    )
    assert (tmp_path / "x.xlsx").exists() and (tmp_path / "x.pdf").exists()


def test_exports_include_rack_devices(tmp_path: Path, catalog: Catalog) -> None:
    from sitesizer.exporters.rack import RackDiagram, rack_table_rows

    key = size_site(make_site(sockets=48), catalog).rack.plans[0].key
    extras = [
        {"id": "d1", "kind": "device", "model": "FS-148F", "rack": key, "u": 2},
        {"id": "c1", "kind": "custom", "label": "NVR", "rack": key, "u": 3},
    ]
    r = size_site(make_site(sockets=48, layout={"extras": extras}), catalog)
    wb = load_workbook(export_xlsx(r, tmp_path / "x.xlsx", catalog=catalog))
    cells = {str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None}
    assert any("FS-148F" in v for v in cells) and any("NVR" in v for v in cells)
    assert any(row and "FS-148F" in " ".join(map(str, row)) for row in rack_table_rows(r))
    assert export_png(RackDiagram(r, catalog), tmp_path / "r.png").exists()  # type: ignore[arg-type]
    assert export_pdf(r, tmp_path / "r.pdf", catalog=catalog).stat().st_size > 1000
