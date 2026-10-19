"""End-to-end GUI smoke test (offscreen): drives the real main window like a user would."""

from __future__ import annotations

import os
from pathlib import Path

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtCore import QPointF, QSettings, Qt, QThreadPool
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication


@pytest.fixture(scope="module")
def app() -> QApplication:
    inst = QApplication.instance() or QApplication([])
    from sitesizer.gui.theme import load_fonts, set_font_family

    set_font_family(load_fonts())
    return inst  # type: ignore[return-value]


@pytest.fixture
def window(app: QApplication, tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    from sitesizer.gui import state as state_mod
    from sitesizer.gui.main_window import MainWindow
    from sitesizer.gui.theme import theme_manager

    monkeypatch.setattr(state_mod, "app_data_dir", lambda: tmp_path)
    monkeypatch.setattr(state_mod, "user_catalog_path", lambda: tmp_path / "catalog.json")
    qs = QSettings(str(tmp_path / "s.ini"), QSettings.Format.IniFormat)
    st = state_mod.AppState(qs)
    st.settings.onboarding_done = True
    theme_manager.apply("light", 1.0)
    win = MainWindow(st)
    win._tour_shown = True
    win.resize(1400, 900)
    win.show()
    app.processEvents()
    yield win
    st._dirty_struct = False
    st.undo.setClean()
    win.close()
    win.deleteLater()
    app.processEvents()


def pump(app: QApplication, n: int = 5) -> None:
    for _ in range(n):
        app.processEvents()
    QTest.qWait(80)


def test_typing_updates_result_and_undo(app: QApplication, window) -> None:
    st = window.state
    loc = window.location
    loc.sockets.field.setFocus()
    loc.sockets.field.selectAll()
    QTest.keyClicks(loc.sockets.field, "100")
    pump(app)
    assert st.site.total_sockets == 100
    assert st.result is not None and st.result.categories["access_switch"].count == 3
    assert window.summary.sw.value.text() == str(st.result.total_switches)
    st.undo.undo()
    pump(app)
    assert st.site.total_sockets == 0
    assert loc.sockets.value() == 0
    st.undo.redo()
    pump(app)
    assert st.site.total_sockets == 100


def test_zones_presets_and_tier(app: QApplication, window) -> None:
    st, loc = window.state, window.location
    loc.add_zone()
    loc.add_zone()
    pump(app)
    assert len(st.site.ap_groups) == 2
    loc._zone_rows[0].qty.setValue(30)
    loc._zone_rows[0].qty.field.valueChanged.emit(30)
    pump(app)
    assert st.site.ap_groups[0].qty == 30
    loc._zone_rows[1].remove.click()
    pump(app)
    assert len(st.site.ap_groups) == 1
    loc.apply_preset("warehouse")
    pump(app)
    assert st.site.cameras == 40 and st.site.tier == 4
    QTest.keyClick(loc.tier, Qt.Key.Key_1)
    pump(app)
    assert st.site.tier == 1
    assert loc.psu_confirm.isVisible()
    yes = [b for b in loc.psu_confirm.findChildren(type(loc.add_zone.__self__.sockets.plus)) if b.text()]
    yes[0].click()
    pump(app)
    assert st.site.redundant_psu is True
    assert not loc.psu_confirm.isVisible()


def test_navigation_and_pages_render(app: QApplication, window) -> None:
    window.location.apply_preset("hq")
    pump(app)
    for key in ("bom", "topology", "ipplan", "power", "compare", "projects", "catalog", "help", "settings", "location"):
        window.navigate(key)
        pump(app, 3)
        assert window.stack.currentWidget() is window.pages[key]
    assert window.bom.proxy.rowCount() > 3
    assert window.ipplan.model.rowCount() >= 4
    assert window.compare.model.rowCount() > 0


def test_sites_save_load(app: QApplication, window, tmp_path: Path) -> None:
    st = window.state
    window.site_action("add", "")
    st.set_field("name", "Філія 2", merge=False)
    window.site_action("duplicate", st.current_id)
    pump(app)
    assert len(st.project.sites) == 3
    path = tmp_path / "p.sizing.json"
    saved = st.save_project(path)
    assert saved.exists() and not st.dirty
    st.new_project()
    assert len(st.project.sites) == 1
    window.open_path(str(path), ask=False)
    pump(app)
    assert len(st.project.sites) == 3
    assert window.sidebar.site_group.buttons().__len__() == 3


def test_theme_and_language_switch(app: QApplication, window) -> None:
    from sitesizer.gui.theme import tokens

    window.toggle_theme()
    pump(app)
    assert tokens().dark
    window.toggle_theme()
    pump(app)
    window.state.update_settings(language="en")
    pump(app)
    assert window.topbar.title.text() == "Location"
    assert window.state.result is not None
    window.state.set_field("sockets", 10, merge=False)
    pump(app)
    assert window.state.result.lines("access_switch")[0].category == "Access switches"
    window.state.update_settings(language="uk")
    pump(app)
    assert window.topbar.title.text() == "Локація"
    window.state.update_settings(ui_scale=1.25)
    pump(app)
    assert window.sidebar.width() in (round(236 * 1.25), round(64 * 1.25))
    window.state.update_settings(ui_scale=1.0)
    pump(app)
    assert window.state.result is not None


def test_exports_in_background(app: QApplication, window, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    window.location.apply_preset("rnd_office")
    pump(app)
    targets = iter([tmp_path / "a.xlsx", tmp_path / "a.pdf", tmp_path / "a.png", tmp_path / "a.svg"])
    monkeypatch.setattr(window, "_ask_path", lambda *_a, **_k: next(targets))
    window.export_xlsx()
    window.export_pdf()
    window.export_diagram("png")
    window.export_diagram("svg")
    QThreadPool.globalInstance().waitForDone(20_000)
    pump(app, 10)
    for name in ("a.xlsx", "a.pdf", "a.png", "a.svg"):
        assert (tmp_path / name).exists(), name


def test_palette_and_tour(app: QApplication, window) -> None:
    cmds = window.commands()
    assert any("PDF" in c.title for c in cmds)
    from sitesizer.gui.widgets.overlays import CommandPalette, TourOverlay, fuzzy_score

    assert fuzzy_score("pdf", "PDF-звіт") == 0
    assert fuzzy_score("ткм", "тема") is None
    assert fuzzy_score("тм", "тема") is not None
    pal = CommandPalette(window, cmds)
    pal.edit.setText("тема")
    assert pal.list.count() >= 1
    pal.deleteLater()
    window.start_tour()
    pump(app)
    overlays = window.centralWidget().findChildren(TourOverlay)
    assert overlays and overlays[0].isVisible()
    tour = overlays[0]
    assert len(tour.steps) == 5
    for _ in range(5):
        QTest.keyClick(tour, Qt.Key.Key_Right)
        pump(app, 2)
    assert window.state.settings.onboarding_done


def test_catalog_edit_applies(app: QApplication, window) -> None:
    from sitesizer.gui.views.bom import COL_TOTAL, LINE_ROLE
    from sitesizer.gui.views.catalog_view import MODEL_COLUMNS

    def column(key: str) -> int:
        return next(i for i, c in enumerate(MODEL_COLUMNS) if c[0] == key)

    cv = window.catalog
    keys = cv.models.keys
    row = keys.index("FS-148F")
    assert cv.models.setData(cv.models.index(row, column("cat.col.price")), "12345")
    assert cv.models.setData(cv.models.index(row, column("cat.col.code")), "000084545")
    assert window.state.catalog.device("FS-148F").price == 12345.0
    window.state.set_field("sockets", 48, merge=False)
    pump(app)
    line = window.state.result.lines("access_switch")[0]
    assert line.unit_price == 12345.0 and line.code == "000084545"
    window.navigate("bom")  # hidden pages refresh when they are shown
    pump(app)
    group = window.bom.proxy.index(0, 0)
    assert window.bom.proxy.index(0, COL_TOTAL, group).data(LINE_ROLE) is not None
    bad = cv.models.setData(cv.models.index(keys.index("FS-148F-FPOE"), column("cat.col.ports")), "-5")
    assert not bad and not cv.error.isHidden()
    window.state.reset_catalog()


def test_catalog_add_model(app: QApplication, window) -> None:
    from sitesizer.gui.views.catalog_view import MODEL_COLUMNS, AddModelDialog, new_model

    cv = window.catalog
    dlg = AddModelDialog(set(cv.data["models"]), "", cv)
    dlg.sku.setText("FS-148F")
    dlg._accept()
    assert dlg.result() != dlg.DialogCode.Accepted and not dlg.problem.isHidden()  # duplicate SKU
    dlg.sku.setText("NVR-16")
    dlg.name.setText("NVR 16 ch")
    dlg.power.setText("45,5")
    dlg.price.setText("12000")
    dlg._accept()
    key, model = dlg.values()
    assert key == "NVR-16" and model["power_base_w"] == 45.5 and model["price"] == 12000.0
    assert cv.insert_model(key, model)
    dev = window.state.catalog.device("NVR-16")
    assert dev.kind == "accessory" and dev.rack_units == 1 and dev.name["uk"] == "NVR 16 ch"
    # every kind gets the sections it needs, so a new switch / FortiGate / AP is valid right away
    for kind in ("switch", "firewall", "ap", "license", "work"):
        assert cv.insert_model(f"NEW-{kind}", new_model(kind))
    # changing the type in the table fills the missing sections instead of failing
    col = next(i for i, c in enumerate(MODEL_COLUMNS) if c[0] == "cat.col.kind")
    assert cv.models.setData(cv.models.index(cv.models.keys.index("NVR-16"), col), "switch")
    assert window.state.catalog.device("NVR-16").ports is not None
    # a model from the catalog can be put into a cabinet and is counted everywhere
    assert cv.insert_model("NVR-32", new_model("accessory", rack_units=2, power_w=60))
    st = window.state
    st.set_field("sockets", 48, merge=False)
    window.navigate("racks")
    pump(app)
    window.racks.add_extra("device", 30, 2, "NVR-32")
    pump(app)
    assert any(line.model == "NVR-32" for line in st.result.lines("rack_device"))
    assert ("NVR-32", 1, 60) in st.result.power.breakdown
    st.undo.undo()
    pump(app)
    window.state.reset_catalog()


def _mouse(widget, kind, pos) -> None:
    from PySide6.QtCore import QEvent, QPointF
    from PySide6.QtGui import QMouseEvent

    types = {
        "press": QEvent.Type.MouseButtonPress,
        "move": QEvent.Type.MouseMove,
        "release": QEvent.Type.MouseButtonRelease,
    }
    buttons = Qt.MouseButton.NoButton if kind == "release" else Qt.MouseButton.LeftButton
    p = QPointF(pos)
    ev = QMouseEvent(
        types[kind], p, widget.mapToGlobal(p), Qt.MouseButton.LeftButton, buttons, Qt.KeyboardModifier.NoModifier
    )
    QApplication.sendEvent(widget, ev)


def test_rack_editor_drag_and_drop(app: QApplication, window) -> None:
    st = window.state
    st.edit("t", lambda d: d.update(sockets=96, aggregation="yes", mode="extended"))
    st.recompute()
    window.navigate("racks")
    pump(app)
    editor = window.racks.editor
    plan = st.result.rack.plans[0]
    rect = editor.item_rect_in_widget("access_switch:2")
    target = editor.unit_point(plan.key, 3)
    assert rect is not None and target is not None
    _mouse(editor, "press", rect.center())
    _mouse(editor, "move", rect.center() + (target - rect.center()) / 2)
    _mouse(editor, "move", target)
    _mouse(editor, "release", target)
    pump(app)
    pos = st.site.layout.positions["access_switch:2"]
    assert (pos.rack, pos.u) == (plan.key, 3)
    # keyboard nudge, extra organizer, new cabinet
    editor.setFocus()
    QTest.keyClick(editor, Qt.Key.Key_Up)
    pump(app)
    assert st.site.layout.positions["access_switch:2"].u == 4
    window.racks.add_extra("manager")
    dev_id = window.racks.add_extra("device", 30, 1, "FS-148F")
    pump(app)
    assert dev_id and any(line.group == "rack_device" for line in st.result.bom)
    st.undo.undo()
    pump(app)
    window.racks.add_rack(24)
    pump(app)
    assert len(st.site.layout.extras) == 1 and st.site.layout.added == ["user-1"]
    assert [p.size_u for p in st.result.rack.plans][-1] == 24
    st.undo.undo()
    st.undo.undo()
    pump(app)
    assert st.site.layout.added == [] and not st.site.layout.extras


def test_rack_editor_swap_devices(app: QApplication, window) -> None:
    """Dropping a device onto another one swaps the two."""
    st = window.state
    st.edit("t", lambda d: d.update(sockets=96, cameras=10))
    st.recompute()
    window.navigate("racks")
    pump(app)
    editor = window.racks.editor

    def unit(item_id: str) -> int:
        return next(it.u for p in st.result.rack.plans for it in p.items if it.id == item_id)

    a, b = "access_switch:1", "camera_switch:1"
    ua, ub = unit(a), unit(b)
    ra, rb = editor.item_rect_in_widget(a), editor.item_rect_in_widget(b)
    assert ra is not None and rb is not None
    _mouse(editor, "press", ra.center())
    _mouse(editor, "move", ra.center() + (rb.center() - ra.center()) / 2)
    _mouse(editor, "move", rb.center())
    assert editor._swap is not None
    _mouse(editor, "release", rb.center())
    pump(app)
    assert (unit(a), unit(b)) == (ub, ua)
    assert not st.result.rack.layout_problems
    st.undo.undo()
    pump(app)
    assert (unit(a), unit(b)) == (ua, ub)


def test_bom_inline_edit_and_custom_line(app: QApplication, window) -> None:
    from PySide6.QtWidgets import QSpinBox

    from sitesizer.gui.views.bom import COL_QTY, LINE_ROLE

    st = window.state
    st.edit("t", lambda d: d.update(sockets=96))
    st.recompute()
    window.navigate("bom")
    pump(app)
    bom = window.bom
    idx = None
    for r in range(bom.proxy.rowCount()):
        g = bom.proxy.index(r, 0)
        for c in range(bom.proxy.rowCount(g)):
            child = bom.proxy.index(c, 0, g)
            if child.data(LINE_ROLE).group == "access_switch":
                idx = child.siblingAtColumn(COL_QTY)
    assert idx is not None
    bom.tree.setCurrentIndex(idx)
    bom.tree.edit(idx)
    pump(app)
    editor = bom.tree.findChild(QSpinBox)
    assert editor is not None
    editor.setValue(4)
    QTest.keyClick(editor, Qt.Key.Key_Return)
    pump(app)
    assert st.site.bom.overrides["access_switch:FS-148F"].qty == 4
    assert st.result.categories["access_switch"].count == 4
    line = st.result.lines("access_switch")[0]
    assert line.manual and line.calc_qty == 2
    bom.commit(line, COL_QTY, 2)  # back to the calculated value → override dropped
    pump(app)
    assert "access_switch:FS-148F" not in st.site.bom.overrides
    st.edit(
        "t",
        lambda d: d["bom"]["custom"].append({"id": "w1", "name": "Монтаж", "qty": 3, "price": 100, "section": "works"}),
    )
    st.recompute()
    pump(app)
    assert st.result.lines("custom")[0].total_price == 300


def test_export_dialog_choices(app: QApplication, window, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from openpyxl import load_workbook

    from sitesizer.gui.widgets import export_dialog

    st = window.state
    st.edit("t", lambda d: d.update(sockets=60, location_code="BO7", location_id=7))
    st.recompute()

    def fake_exec(self) -> int:
        self.sheets["prices"].setChecked(False)
        self.sheets["inputs"].setChecked(False)
        self.pdf.setChecked(True)
        self.folder.setText(str(tmp_path))
        self.name.setText("out")
        return 1

    monkeypatch.setattr(export_dialog.ExportDialog, "exec", fake_exec)
    window.export_all()
    QThreadPool.globalInstance().waitForDone(20_000)
    pump(app, 10)
    # each table in its own file by default
    assert sorted(p.name for p in tmp_path.glob("*.xlsx")) == [
        "out — IP-таблиця.xlsx",
        "out — специфікація.xlsx",
        "out — схема шаф.xlsx",
    ]
    assert load_workbook(tmp_path / "out — схема шаф.xlsx").sheetnames == ["BO7 - Шафи"]
    assert (tmp_path / "out.pdf").read_bytes().startswith(b"%PDF")
    assert st.settings.export_choices["only_used"] is True

    def one_file(self) -> int:
        fake_exec(self)
        self.split.setChecked(False)
        return 1

    monkeypatch.setattr(export_dialog.ExportDialog, "exec", one_file)
    window.export_all()
    QThreadPool.globalInstance().waitForDone(20_000)
    pump(app, 10)
    wb = load_workbook(tmp_path / "out.xlsx")
    assert wb.sheetnames == ["Слаботрумка", "BO7 - Схема+шафи", "BO7 - IP"]


def test_settings_and_catalog_move_over_from_the_old_name(tmp_path: Path) -> None:
    from sitesizer.gui.app import migrate_legacy_data

    old_dir = tmp_path / "SiteSizer" / "SiteSizer"
    old_dir.mkdir(parents=True)
    (old_dir / "catalog.json").write_text('{"x": 1}', encoding="utf-8")
    new_dir = tmp_path / "LocalCount" / "LocalCount"
    new_dir.mkdir(parents=True)
    legacy = QSettings(str(tmp_path / "old.ini"), QSettings.Format.IniFormat)
    legacy.setValue("app/settings_json", '{"theme": "dark"}')
    settings = QSettings(str(tmp_path / "new.ini"), QSettings.Format.IniFormat)
    assert migrate_legacy_data(settings, legacy, new_dir)
    assert settings.value("app/settings_json") == '{"theme": "dark"}'
    assert (new_dir / "catalog.json").read_text(encoding="utf-8") == '{"x": 1}'
    # the second start leaves the new data alone
    (new_dir / "catalog.json").write_text('{"x": 2}', encoding="utf-8")
    settings.setValue("app/settings_json", '{"theme": "light"}')
    assert not migrate_legacy_data(settings, legacy, new_dir)
    assert (new_dir / "catalog.json").read_text(encoding="utf-8") == '{"x": 2}'
    assert settings.value("app/settings_json") == '{"theme": "light"}'


def test_racks_rooms_rename_and_delete(app: QApplication, window, monkeypatch: pytest.MonkeyPatch) -> None:
    from sitesizer.gui.views import racks as racks_mod

    monkeypatch.setattr(racks_mod, "confirm", lambda *a, **k: True)
    st = window.state
    racks = window.racks

    def settle() -> None:
        pump(app)
        st.recompute()
        pump(app)

    def chip_texts() -> list[str]:
        return [racks.chips.itemAt(i).widget().text() for i in range(racks.chips.count())]

    st.edit("t", lambda d: d.update(sockets=96, mode="extended"))
    st.recompute()
    window.navigate("racks")
    pump(app)
    # one room: "all cabinets", the room and "+ room"
    assert chip_texts() == ["Усі шафи (1)", "Комутаційна 1 (MDF) · 1 шафа", " Комутаційна"]
    # add a room: the switches are spread over it and the view jumps to it
    racks.add_room_btn.click()
    settle()
    assert st.site.closets == 2 and st.result.rack.rooms[1] == "Комутаційна 2 (IDF-1)"
    assert racks.view_room == 1
    assert [p.room for p in racks.editor.diagram.plans] == [1]
    window.navigate("location")
    pump(app)
    assert window.location.closets.value() == 2
    window.navigate("racks")
    pump(app)
    # add two cabinets to that room, then look at it and at everything
    racks.add_rack(42)
    settle()
    racks.add_rack(24)
    settle()
    assert [p.room for p in st.result.rack.plans] == [0, 1, 1, 1]
    assert len(racks.editor.diagram.plans) == 3 and racks.view_room == 1
    assert "3 шафи" in chip_texts()[2]
    racks.chips.itemAt(0).widget().click()  # "all cabinets"
    pump(app)
    assert racks.view_room is None and len(racks.editor.diagram.plans) == 4
    racks.chips.itemAt(1).widget().click()  # the main room
    pump(app)
    assert [p.room for p in racks.editor.diagram.plans] == [0] and racks.chips.itemAt(1).widget().isChecked()
    assert not racks.del_room.isEnabled()  # the main room stays
    # delete a cabinet of room 2
    racks.show_room(1)
    pump(app)
    racks.editor.select("user-2")
    racks.delete_rack()
    settle()
    assert [p.key for p in st.result.rack.plans if p.room == 1] == ["idf1-1", "user-1"]
    # rename the room; empty name = automatic again
    racks.room_name.setText("Серверна 2 поверх")
    racks.room_name.editingFinished.emit()
    settle()
    assert st.result.rack.rooms[1] == "Серверна 2 поверх" and "Серверна 2 поверх · 2 шафи" in chip_texts()
    # a third room, then remove the second: the third one takes its place
    racks.add_room()
    settle()
    assert len(st.result.rack.rooms) == 3
    racks.delete_room(1)
    settle()
    assert st.site.closets == 2 and st.result.rack.rooms == ["Комутаційна 1 (MDF)", "Комутаційна 2 (IDF-1)"]
    assert not st.site.layout.added and not st.site.layout.room_names
    st.undo.undo()
    settle()
    assert st.result.rack.rooms[1] == "Серверна 2 поверх" and st.site.layout.added == ["user-1"]
    # a room without cabinets shows a hint instead of an empty canvas
    racks.show_room(2)
    for p in [p for p in st.result.rack.plans if p.room == 2]:
        racks.editor.sel_rack = p.key
        racks.delete_rack()
        settle()
    assert racks.view_room == 2 and racks.empty.isVisible() and not racks.scroll.isVisible()
    racks.add_rack(24)
    settle()
    assert [p.key for p in st.result.rack.plans if p.room == 2] and racks.scroll.isVisible()


def test_rack_items_rename_and_delete(app: QApplication, window) -> None:
    st = window.state
    racks = window.racks
    st.edit("t", lambda d: d.update(sockets=96, mode="extended"))
    st.recompute()
    window.navigate("racks")
    pump(app)
    # name a switch; the model stays visible under the name
    plan = st.result.rack.plans[0]
    sw = next(it for it in plan.items if it.group == "access_switch")
    racks.editor.select(plan.key, sw.id)
    pump(app)
    racks.item_name.setText("Комутатор холу")
    racks.item_name.editingFinished.emit()
    st.recompute()
    pump(app)
    assert st.site.layout.labels == {sw.id: "Комутатор холу"}
    assert racks.editor.item(sw.id).label == "Комутатор холу"
    assert racks.item_title.text() == "Комутатор холу" and racks.item_model.text() == sw.model
    racks.item_name.setText("")
    racks.item_name.editingFinished.emit()
    st.recompute()
    pump(app)
    assert st.site.layout.labels == {} and racks.editor.item(sw.id).label == sw.label
    # an added device can be renamed and deleted (it leaves the bill of materials too)
    racks.editor.select(plan.key)
    dev_id = racks.add_extra("device", 30, 1, "FS-148F", "NVR")
    st.recompute()
    pump(app)
    assert racks.editor.sel_item == dev_id and racks.del_item.isVisible()
    racks.item_name.setText("Відеосервер")
    racks.item_name.editingFinished.emit()
    st.recompute()
    pump(app)
    # a switch put in by hand is a network device: its rename is kept like that of an automatic one
    assert st.site.layout.labels[dev_id] == "Відеосервер" and racks.editor.item(dev_id).label == "Відеосервер"
    assert any(line.group == "rack_device" for line in st.result.bom)
    racks.del_item.click()
    st.recompute()
    pump(app)
    assert not st.site.layout.extras and racks.editor.item(dev_id) is None
    assert not any(line.group == "rack_device" for line in st.result.bom)
    # active equipment can be deleted too — the firewall and switches leave the specification
    fw = next(it for it in st.result.rack.plans[0].items if it.group == "firewall")
    racks.editor.select(plan.key, fw.id)
    racks.editor.select(plan.key, sw.id, add=True)  # Ctrl+click
    pump(app)
    assert racks.editor.sel_items == [fw.id, sw.id] and racks.del_item.isVisible()
    racks.del_item.click()
    st.recompute()
    pump(app)
    assert not st.result.lines("firewall")
    assert st.result.lines("access_switch")[0].qty == 1
    assert racks.restore_btn.isVisible()
    racks.restore_deleted()
    st.recompute()
    pump(app)
    assert st.result.lines("firewall") and st.result.lines("access_switch")[0].qty == 2


def test_rack_multi_select_drag_and_copy(app: QApplication, window) -> None:
    st = window.state
    racks = window.racks
    st.edit("t", lambda d: d.update(sockets=96, location_code="BO123"))
    st.recompute()
    window.navigate("racks")
    pump(app)
    editor = racks.editor
    plan = st.result.rack.plans[0]
    sws = [it for it in plan.items if it.group == "access_switch"]
    editor.select(plan.key, sws[0].id)
    editor.select(plan.key, sws[1].id, add=True)
    pump(app)
    # Ctrl+C copies the names of the selected devices, one per line
    editor.setFocus()
    QTest.keyClick(editor, Qt.Key.Key_C, Qt.KeyboardModifier.ControlModifier)
    pump(app)
    assert QApplication.clipboard().text() == "BO123-1A0-ASW01\nBO123-1A0-ASW02"
    # the "copy names" button copies every device of the cabinets shown
    editor.select(plan.key)
    racks.copy_names()
    assert QApplication.clipboard().text().splitlines()[0] == "BO123-1A0-FW01"
    # dragging one of two selected devices moves both by the same distance
    editor.select(plan.key, sws[0].id)
    editor.select(plan.key, sws[1].id, add=True)
    rect = editor.item_rect_in_widget(sws[1].id)
    target = editor.unit_point(plan.key, 3)
    assert rect is not None and target is not None
    _mouse(editor, "press", rect.center())
    _mouse(editor, "move", rect.center() + (target - rect.center()) / 2)
    _mouse(editor, "move", target)
    _mouse(editor, "release", target)
    pump(app)
    pos = st.site.layout.positions
    delta = 3 - sws[1].u
    assert (pos[sws[0].id].u, pos[sws[1].id].u) == (sws[0].u + delta, 3)
    st.undo.undo()
    pump(app)
    assert not st.site.layout.positions


def test_rack_cut_and_paste_into_another_cabinet(app: QApplication, window) -> None:
    st = window.state
    racks = window.racks
    st.edit("t", lambda d: d.update(sockets=96, location_code="BO123"))
    st.recompute()
    window.navigate("racks")
    pump(app)
    racks.add_rack(42)
    st.recompute()
    pump(app)
    editor = racks.editor
    src, dst = st.result.rack.plans[0], st.result.rack.plans[-1]
    assert src.key != dst.key
    sws = [it for it in src.items if it.group == "access_switch"]
    editor.select(src.key, sws[0].id)
    editor.select(src.key, sws[1].id, add=True)
    editor.setFocus()
    QTest.keyClick(editor, Qt.Key.Key_X, Qt.KeyboardModifier.ControlModifier)
    assert editor.cut_ids == [sws[0].id, sws[1].id]
    # click unit 30 of the other cabinet, paste: both move there, spaced as before, top at U30
    pt = editor.unit_point(dst.key, 30)
    assert pt is not None
    _mouse(editor, "press", pt)
    _mouse(editor, "release", pt)
    QTest.keyClick(editor, Qt.Key.Key_V, Qt.KeyboardModifier.ControlModifier)
    pump(app)
    pos = st.site.layout.positions
    assert pos[sws[0].id].rack == dst.key and pos[sws[1].id].rack == dst.key
    assert pos[sws[0].id].u == 30 and pos[sws[0].id].u - pos[sws[1].id].u == sws[0].u - sws[1].u
    assert not editor.cut_ids
    st.undo.undo()
    pump(app)
    assert not st.site.layout.positions


def test_floors_of_one_location(app: QApplication, window) -> None:
    st = window.state
    st.edit("t", lambda d: d.update(sockets=96, location_code="BO123"))
    first = st.current_id
    new_id = st.add_site()
    pump(app)
    assert st.site.floor == 2 and st.site.name == "2 поверх" and st.site.location_code == "BO123"
    st.edit("t", lambda d: d.update(sockets=48, location_id=57))
    st.recompute()
    pump(app)
    # location-wide fields follow on every floor; one undo step reverts both
    assert st.project.site(first).input.location_id == 57
    # one firewall for the location, on the first floor; this floor links to it with an ODF
    assert not st.result.lines("firewall")
    assert any(it.label == "ODF 1A0" for p in st.result.rack.plans for it in p.items)
    combined = st.location_result
    assert combined is not None and combined.lines("firewall")[0].qty == 1
    assert combined.lines("access_switch")[0].qty == 3
    window.navigate("bom")
    pump(app)
    assert window.bom.scope_ctl.isVisible() and window.bom.shown() is combined
    # move the firewall to this floor
    window.navigate("location")
    pump(app)
    window.location.fw_here.switch.setChecked(True)
    pump(app)
    assert st.project.site(new_id).is_hub and st.result.lines("firewall")
    st.undo.undo()
    pump(app)
    assert st.project.site(first).input.location_id is None


def test_location_id_typed_and_cleared(app: QApplication, window) -> None:
    st = window.state
    field = window.location.loc_id
    assert field.text() == "" and st.site.location_id is None
    field.setFocus()
    QTest.keyClicks(field, "57")
    pump(app)
    assert st.site.location_id == 57
    QTest.keyClicks(field, "9")  # 579 > 255 is refused
    pump(app)
    assert field.text() == "57" and st.site.location_id == 57
    field.clear()
    pump(app)
    assert st.site.location_id is None
    field.clearFocus()
    QTest.keyClicks(field, "12")
    pump(app)
    assert st.site.location_id == 12
    field.clearFocus()
    while st.undo.canUndo():
        st.undo.undo()
    pump(app)
    assert st.site.location_id is None and field.text() == ""


def test_topology_editor(app: QApplication, window) -> None:
    st = window.state
    st.edit("t", lambda d: d.update(sockets=240, cameras=30, tier=1, aggregation="yes", location_code="BO123"))
    st.recompute()
    window.navigate("topology")
    pump(app)
    view = window.topology
    canvas = view.canvas
    d = canvas.diagram
    assert d is not None
    sw = next(n for n in d.topo.nodes if n.kind == "access_switch")
    # dragging a device stores its place (one undoable edit, shared by the location)
    start = canvas.node_point(sw.id)
    assert start is not None
    target = start + QPointF(60, 40)
    _mouse(canvas.viewport(), "press", start)
    _mouse(canvas.viewport(), "move", start + QPointF(30, 20))
    _mouse(canvas.viewport(), "move", target)
    _mouse(canvas.viewport(), "release", target)
    pump(app)
    assert sw.id in st.site.topology.positions
    # remove a link, draw a new one, change the access design
    stack = next(lk for lk in canvas.diagram.topo.links if lk.kind == "stack")
    view.delete_links([stack.id])
    st.recompute()
    pump(app)
    assert stack.id in st.site.topology.hidden and canvas.diagram.topo.link(stack.id) is None
    fw = next(n for n in canvas.diagram.topo.nodes if n.kind == "firewall")
    canvas.select("node", sw.id)
    canvas.select("node", fw.id, add=True)
    view._connect_selected()
    st.recompute()
    pump(app)
    assert any(lk.a == sw.id and lk.b == fw.id and not lk.auto for lk in canvas.diagram.topo.links)
    view.access.setValue("dual")
    view.access.valueChanged.emit("dual")
    st.recompute()
    pump(app)
    assert st.site.topology.access == "dual"
    assert not any(lk.kind == "stack" for lk in canvas.diagram.topo.links)
    view.reset()
    st.recompute()
    pump(app)
    assert st.site.topology.is_empty
    st.undo.undo()  # the reset is undoable
    assert st.site.topology.access == "dual" and st.site.topology.links
    for _ in range(4):
        st.undo.undo()
    st.recompute()
    pump(app)
    assert st.site.topology.is_empty and st.site.topology.access == "ring"


def test_location_equipment_choices(app: QApplication, window, tmp_path: Path) -> None:
    from sitesizer.gui.views.catalog_view import new_model

    st = window.state
    loc = window.location
    window.navigate("location")
    st.set_field("sockets", 96, merge=False)
    pump(app)
    # firewall model and number
    fws = [m for m, d in st.catalog.models.items() if d.kind == "firewall"]
    other = next(m for m in fws if m != st.result.firewall.model)
    loc.fw_model.setCurrentIndex(loc.fw_model.findData(other))
    pump(app)
    assert st.site.firewall_model == other and st.result.firewall.model == other
    loc.fw_count.changed.emit(2)
    pump(app)
    assert st.result.firewall.count == 2 and not loc.fw_count.reset.isHidden()
    loc.fw_count.changed.emit(None)
    pump(app)
    assert st.site.firewall_count is None
    # core switches typed
    loc.core_count.changed.emit(2)
    pump(app)
    assert st.result.core.count == 2
    # switch count typed in the room block, then back to the calculation
    block = loc._room_blocks[0]
    block.counts["access_switch"].changed.emit(4)
    pump(app)
    assert st.result.categories["access_switch"].count == 4
    assert block.counts["access_switch"].stepper.value() == 4
    block.counts["access_switch"].changed.emit(None)
    pump(app)
    assert st.result.categories["access_switch"].count == 2
    # UPS of the room's cabinet
    ups = [m for m, _d in st.catalog.ups_models()]
    block.ups_model.setCurrentIndex(block.ups_model.findData(ups[0]))
    block.ups_count.changed.emit(1)
    pump(app)
    assert [it.model for p in st.result.rack.plans for it in p.items if it.id.startswith("ups:")] == [ups[0]]
    # a further device: in the cabinet, the specification and the topology
    block._add_device()
    pump(app)
    assert len(st.site.rooms[0].devices) == 1 and len(block.dev_rows) == 1
    row = block.dev_rows[0]
    row.model.setCurrentIndex(row.model.findData("FS-148F"))
    pump(app)
    assert st.site.rooms[0].devices[0].model == "FS-148F"
    assert any(it.id.startswith("dev:") for p in st.result.rack.plans for it in p.items)
    row.remove.click()
    pump(app)
    assert not st.site.rooms[0].devices
    # a switch added to the catalog is offered for the room right away
    assert window.catalog.insert_model("MY-SW-48", new_model("switch"))
    pump(app)
    block = loc._room_blocks[0]
    assert block.combos["access_switch"].findData("MY-SW-48") > 0
    assert loc.core_model.findData("MY-SW-48") > 0
    window.state.reset_catalog()
    pump(app)


def test_ap_wiring_dialog(app: QApplication, window, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from sitesizer.gui.widgets.ap_wiring_dialog import ApWiringDialog, location_wiring

    st = window.state
    st.edit("zones", lambda d: d.update(ap_groups=[{"zone": "low_density", "qty": 40}]))
    pump(app)
    w = location_wiring(st)
    assert len(w.rooms[0].aps) == 40 and w.rooms[0].shared_pairs == 0
    dlg = ApWiringDialog(st, window.topology.export_requested.emit, window)
    dlg.columns.setValue(5)
    pump(app)
    assert st.site.ap_columns == 5
    out = tmp_path / "ap.xlsx"
    monkeypatch.setattr(window, "_ask_path", lambda *a, **k: out)
    window.export_diagram("apw_xlsx")
    assert out.exists()
    out_png = tmp_path / "ap.png"
    monkeypatch.setattr(window, "_ask_path", lambda *a, **k: out_png)
    window.export_diagram("apw_png")
    assert out_png.exists()
    dlg.accept()


def test_wifi_switches_follow_the_aps(app: QApplication, window) -> None:
    st = window.state
    loc = window.location
    st.edit("zones", lambda d: d.update(ap_groups=[{"zone": "low_density", "qty": 12}]))
    pump(app)
    loc._room_blocks[0].counts["wifi_switch"].changed.emit(1)  # typed by hand
    pump(app)
    assert st.site.rooms[0].counts == {"wifi_switch": 1}
    row = loc._zone_rows[0]
    row.qty.setValue(122)
    row._emit_qty(122)
    pump(app)
    assert st.site.rooms[0].counts == {}
    assert st.result.categories["wifi_switch"].count == 6
    # sockets: the access switches follow too
    loc._room_blocks[0].counts["access_switch"].changed.emit(4)
    pump(app)
    loc._room_blocks[0].fields["sockets"].valueChanged.emit(200)
    pump(app)
    assert "access_switch" not in st.site.rooms[0].counts


def test_topology_connection_panel(app: QApplication, window) -> None:
    st = window.state
    st.edit("t", lambda d: d.update(sockets=240, cameras=30, tier=1, aggregation="yes", location_code="BO123"))
    st.recompute()
    window.navigate("topology")
    pump(app)
    view = window.topology
    panel = view.panel
    assert not panel.hint.isHidden() and panel.link_box.isHidden()  # nothing selected: the hint
    topo = view.canvas.diagram.topo
    stack = next(lk for lk in topo.links if lk.kind == "stack")
    view.canvas.select("link", stack.id)
    pump(app)
    assert not panel.link_box.isHidden() and panel.link is not None and panel.link.id == stack.id
    port_a = panel.ends["a"][1]
    assert port_a.text() == "" and port_a.placeholderText() == stack.port_a  # automatic
    assert panel.medium.value() == "auto"
    # modules + fibre: a transceiver at both ends
    panel.medium.setValue("fiber")
    panel.medium.valueChanged.emit("fiber")
    st.recompute()
    pump(app)
    lk = view.canvas.diagram.topo.link(stack.id)
    assert st.site.topology.media == {stack.id: "fiber"} and lk.module_a and lk.module_b
    # typed port, module and cable
    view.canvas.select("link", stack.id)
    pump(app)
    port_a = panel.ends["a"][1]
    port_a.setText("sfp+3")
    panel.ends["b"][2].setCurrentText("SFP Foxgate 1330nm")
    panel.cable.setCurrentText("LC-LC OM4 2 м")
    panel._emit()
    st.recompute()
    pump(app)
    lk = view.canvas.diagram.topo.link(stack.id)
    assert lk.port_a == "sfp+3" and lk.module_b == "SFP Foxgate 1330nm" and lk.label == "LC-LC OM4 2 м"
    assert stack.end(stack.b) not in st.site.topology.ports  # an untouched field stays automatic
    # everything automatic again
    view.canvas.select("link", stack.id)
    pump(app)
    panel._reset()
    st.recompute()
    pump(app)
    ed = st.site.topology
    assert not ed.ports and not ed.modules and not ed.media and stack.id not in ed.labels
    # a switch: what its ports serve
    sw = next(n for n in view.canvas.diagram.topo.nodes if n.kind == "access_switch")
    view.canvas.select("node", sw.id)
    pump(app)
    assert not panel.node_box.isHidden() and panel.endpoints.placeholderText().endswith("Розетки")
    panel.endpoints.setText("1–40 Розетки · 48 UPS")
    panel._emit_endpoints()
    st.recompute()
    pump(app)
    assert view.canvas.diagram.topo.node(sw.id).endpoints == "1–40 Розетки · 48 UPS"
    # double-click on a link: the panel's first port field
    view.rename("link", stack.id)
    pump(app)
    assert panel.link is not None and panel.link.id == stack.id


def test_rack_item_model_change(app: QApplication, window) -> None:
    """Clicking a device in a cabinet: another model, same name and place."""
    st = window.state
    st.edit("t", lambda d: d.update(sockets=96))
    st.recompute()
    window.navigate("racks")
    pump(app)
    racks = window.racks

    def item(item_id: str):
        return next(it for p in st.result.rack.plans for it in p.items if it.id == item_id)

    before = item("access_switch:2")
    racks.editor.sel_item = "access_switch:2"
    racks.editor.sel_items = {"access_switch:2"}
    racks._load_side()
    assert racks.model_row.isVisibleTo(racks)
    box = racks.item_model_box
    idx = box.findData("FS-348G")
    assert idx > 0 and box.findData("FAP-231K") < 0  # switches only
    box.setCurrentIndex(idx)
    racks.change_model("access_switch:2")
    st.recompute()
    pump(app)
    after = item("access_switch:2")
    assert after.model == "FS-348G" and after.label == before.label and after.u == before.u
    assert any(line.model == "FS-348G" for line in st.result.bom)
    st.undo.undo()
    st.recompute()
    pump(app)
    assert item("access_switch:2").model == before.model


def test_room_cabinets_and_device_ids(app: QApplication, window) -> None:
    """The number of cabinets of a telecom room is typed on the location page or the racks page."""
    st = window.state
    st.edit("t", lambda d: d.update(closets=2, rooms=[{"sockets": 96}, {"sockets": 48}], location_code="BO1"))
    st.recompute()
    window.navigate("location")
    pump(app)
    block = window.location._room_blocks[1]
    block.cabinets.valueChanged.emit(2)  # typed by the user
    st.recompute()
    pump(app)
    assert st.site.rooms[1].cabinets == 2
    assert [p.tag for p in st.result.rack.plans if p.room == 1] == ["1B0", "1B1"]
    # racks page: the room card shows and changes it
    window.navigate("racks")
    pump(app)
    racks = window.racks
    racks.show_room(1)
    pump(app)
    assert racks.room_cabinets.value() == 2
    racks.room_cabinets.setValue(3)
    st.recompute()
    pump(app)
    assert [p.tag for p in st.result.rack.plans if p.room == 1] == ["1B0", "1B1", "1B2"]
    # a device added to a room gets a stable id: removing another one never renames its cabinet item
    window.navigate("location")
    pump(app)
    block = window.location._room_blocks[1]
    block._add_device()
    block._add_device()
    st.recompute()
    pump(app)
    ids = [d.id for d in st.site.rooms[1].devices]
    assert len(ids) == 2 and all(ids) and ids[0] != ids[1]
    block._remove_device(block.dev_rows[0])
    st.recompute()
    pump(app)
    assert [d.id for d in st.site.rooms[1].devices] == ids[1:]


def test_topology_odf_panel_and_wizard(app: QApplication, window) -> None:
    st = window.state
    st.edit(
        "t",
        lambda d: d.update(
            tier=1,
            aggregation="yes",
            location_code="BO72",
            rooms=[{"sockets": 120, "cameras": 20}, {"sockets": 96, "cameras": 10}],
        ),
    )
    st.recompute()
    window.navigate("topology")
    pump(app)
    view = window.topology
    panel = view.panel
    assert view.odf.isChecked() and not view.panels.isChecked()
    topo = view.canvas.diagram.topo
    odf = next(n for n in topo.nodes if n.kind == "odf" and not n.head and len(n.cells) >= 2)
    # a click on the ODF: its ports in the side panel, a row per duplex with the link in it
    view.canvas.select("node", odf.id)
    pump(app)
    assert not panel.pas_box.isHidden() and panel.link_box.isHidden()
    assert panel.table.rowCount() == odf.ports
    assert panel.table.cellWidget(0, 1).currentData() == odf.cells[1]
    # put the second link into port 1: the two swap (the far end of the cable follows)
    first, second = odf.cells[1], odf.cells[2]
    combo = panel.table.cellWidget(0, 1)
    combo.setCurrentIndex(combo.findData(second))
    combo.activated.emit(combo.currentIndex())
    st.recompute()
    pump(app)
    o = view.canvas.diagram.topo.node(odf.id)
    assert o.cells[1] == second and o.cells[2] == first
    assert view.canvas.diagram.topo.node(o.peer).cells[1] == second
    # a note on a port, the duplexes and the cable
    view.canvas.select("node", odf.id)
    pump(app)
    panel.table.item(4, 2).setText("резерв: КПП")
    panel.pas_ports.setValue(4)
    panel.pas_cable.setText("ВОК OS2 8 вол.")
    panel._emit_passive()
    st.recompute()
    pump(app)
    o = view.canvas.diagram.topo.node(odf.id)
    assert o.notes == {5: "резерв: КПП"} and o.ports == 4 and o.cable == "ВОК OS2 8 вол."
    assert st.site.topology.odf_ports == {odf.id: 4}
    # the ODF's ports automatic again
    view.set_passive(odf.id, {"reset": True})
    st.recompute()
    pump(app)
    assert not st.site.topology.odf_map
    # copper patch panels on: a click shows the switch ports of the panel's ports
    view.panels.setChecked(True)
    st.recompute()
    pump(app)
    pp = next(n for n in view.canvas.diagram.topo.nodes if n.kind == "panel")
    view.canvas.select("node", pp.id)
    pump(app)
    assert not panel.pas_box.isHidden() and panel.table.rowCount() == pp.ports
    assert panel.table.item(0, 1).text() == "1"
    # the wizard: answers on an ODF page and a link page become one undoable edit
    from sitesizer.gui.widgets.connection_wizard import ConnectionWizard

    topo = view.canvas.diagram.topo
    dlg = ConnectionWizard(topo, view.edits, st.catalog, view)
    assert dlg.scopes["odf"][0].isChecked() and dlg.scopes["fibre"][0].isChecked()
    dlg._go(1)  # the first ODF
    item = dlg.items[dlg.at]
    assert item.kind == "odf"
    dlg._fields["name"].setText("ODF провайдера")
    dlg._go(1)
    while dlg.items[dlg.at].kind != "link":
        dlg._skip()
    lk = topo.link(dlg.items[dlg.at].id)
    dlg._fields["cable"].setText("LC-LC OS2 3 м")
    end_a = lk.end(lk.a)
    dlg._fields["ends"][end_a][0].setText("sfp+3")
    dlg._finish()
    assert dlg.result() and dlg.changes
    before = st.site.topology.model_copy(deep=True)
    view._edit("wizard", lambda t: dlg.changes.apply(t))
    st.recompute()
    pump(app)
    topo = view.canvas.diagram.topo
    assert topo.node(item.id).label == "ODF провайдера"
    assert topo.link(lk.id).label == "LC-LC OS2 3 м" and topo.link(lk.id).port_a == "sfp+3"
    st.undo.undo()
    pump(app)
    assert st.site.topology == before
