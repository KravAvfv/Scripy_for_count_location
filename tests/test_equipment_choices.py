"""Equipment chosen on the location page, devices / UPS per telecom room, hand-placed network
devices in the topology, and the AP wiring scheme."""

from __future__ import annotations

from sitesizer.core.ap_wiring import assign, plan_ap_wiring
from sitesizer.core.catalog import Catalog
from sitesizer.core.sizing import size_site
from sitesizer.core.topology import build_topology

from .conftest import make_site


def _items(r, group: str) -> list:
    return [it for p in r.rack.plans for it in p.items if it.group == group]


def _firewalls(catalog: Catalog) -> list[str]:
    return [m for m, d in catalog.models.items() if d.kind == "firewall"]


def test_firewall_model_and_count_chosen_by_user(catalog: Catalog) -> None:
    auto = size_site(make_site(sockets=96), catalog)
    other = next(m for m in _firewalls(catalog) if m != auto.firewall.model)
    r = size_site(make_site(sockets=96, firewall_model=other, firewall_count=2), catalog)
    assert r.firewall.model == other and r.firewall.count == 2
    fws = _items(r, "firewall")
    assert [it.model for it in fws] == [other, other]
    assert {line.model: line.qty for line in r.lines("firewall")} == {other: 2}
    # none at all
    r0 = size_site(make_site(sockets=96, firewall_count=0), catalog)
    assert not _items(r0, "firewall") and not r0.lines("firewall")
    # an unknown model falls back to the calculated one
    bad = size_site(make_site(sockets=96, firewall_model="NOPE"), catalog)
    assert bad.firewall.model == auto.firewall.model


def test_core_count_typed(catalog: Catalog) -> None:
    r = size_site(make_site(sockets=96, core_count=2), catalog)
    assert r.core.count == 2 and len(_items(r, "core_switch")) == 2
    assert r.core.model == catalog.categories["core_switch"].base
    r0 = size_site(make_site(sockets=96, aggregation="yes", core_count=0), catalog)
    assert r0.core.count == 0 and not _items(r0, "core_switch")
    sw = next(m for m, d in catalog.models.items() if d.kind == "switch" and m != r.core.model)
    r2 = size_site(make_site(sockets=96, core_count=1, core_model=sw), catalog)
    assert [it.model for it in _items(r2, "core_switch")] == [sw]


def test_switch_count_typed_per_room(catalog: Catalog) -> None:
    rooms = [{"sockets": 96, "counts": {"access_switch": 5}}, {"sockets": 48}]
    r = size_site(make_site(rooms=rooms), catalog)
    res = r.categories["access_switch"]
    assert res.per_room == [5, 1] and res.count == 6
    assert sum(1 for p in r.rack.plans if p.room == 0 for it in p.items if it.group == "access_switch") == 5
    assert any(c.code == "ROOM_COUNT" for c in r.checks)
    # switches without endpoints: a typed count is still placed
    r2 = size_site(make_site(rooms=[{"counts": {"camera_switch": 2}}]), catalog)
    assert r2.categories["camera_switch"].count == 2 and len(_items(r2, "camera_switch")) == 2


def test_room_devices_and_ups(catalog: Catalog) -> None:
    ups = [m for m, d in catalog.models.items() if d.ups_va]
    rooms = [
        {"sockets": 48, "ups_model": ups[0], "ups_count": 2},
        {"sockets": 48, "ups_model": ups[-1], "ups_count": 1, "devices": [{"model": "FS-148F", "qty": 2}]},
    ]
    r = size_site(make_site(rooms=rooms, location_code="BO1"), catalog)
    by_room: dict[int, list[str]] = {}
    for p in r.rack.plans:
        for it in p.items:
            if it.id.startswith("ups:"):
                by_room.setdefault(p.room, []).append(it.model)
    assert by_room == {0: [ups[0]] * 2, 1: [ups[-1]]}
    ups_lines = {line.model: line.qty for line in r.lines("power") if line.model in ups}
    expected: dict[str, int] = {}
    for m in [ups[0], ups[0], ups[-1]]:
        expected[m] = expected.get(m, 0) + 1
    assert ups_lines == expected
    devs = [it for p in r.rack.plans if p.room == 1 for it in p.items if it.id.startswith("dev:")]
    assert len(devs) == 2 and all(it.group == "access_switch" and "ASW" in it.label for it in devs)
    assert {line.model: line.qty for line in r.lines("rack_device")} == {"FS-148F": 2}
    # no UPS anywhere when the user says 0
    r0 = size_site(make_site(rooms=[{"sockets": 48, "ups_count": 0}], tier=4), catalog)
    assert not [it for p in r0.rack.plans for it in p.items if it.id.startswith("ups:")]


def test_hand_placed_firewall_and_core_appear_in_topology(catalog: Catalog) -> None:
    base = size_site(make_site(sockets=96, location_code="BO115", firewall_count=0, core_count=1), catalog)
    key = base.rack.plans[0].key
    fw = _firewalls(catalog)[0]
    core = catalog.categories["core_switch"].base
    extras = [
        {"id": "x1", "kind": "device", "model": fw, "rack": key, "u": 40, "label": "BO115-1A-CSW01"},
        {"id": "x2", "kind": "device", "model": fw, "rack": key, "u": 38},
        {"id": "x3", "kind": "device", "model": core, "rack": key, "u": 30, "label": "BO115-1A-DSW02"},
    ]
    r = size_site(
        make_site(sockets=96, location_code="BO115", firewall_count=0, core_count=1, layout={"extras": extras}),
        catalog,
    )
    items = {it.id: it for p in r.rack.plans for it in p.items}
    assert items["x1"].group == "firewall" and items["x3"].group == "core_switch"
    assert items["x2"].label.startswith("BO115-1A-FW")  # unnamed: automatic name
    topo = build_topology(r, catalog)
    kinds = {n.label: n.kind for n in topo.nodes}
    assert kinds["BO115-1A-CSW01"] == "firewall" and kinds["BO115-1A-DSW02"] == "core_switch"
    assert sum(1 for n in topo.nodes if n.kind == "core_switch") == 2
    assert any(lk.kind == "isl" for lk in topo.links) and any(lk.kind == "ha" for lk in topo.links)
    # in the specification once, as hand-placed devices
    assert {line.model: line.qty for line in r.lines("rack_device")} == {fw: 2, core: 1}


def test_cabinet_moved_to_another_room(catalog: Catalog) -> None:
    site = make_site(rooms=[{"sockets": 300}, {"sockets": 48}])
    r = size_site(site, catalog)
    second = [p for p in r.rack.plans if p.room == 0][-1]
    moved = size_site(
        make_site(rooms=[{"sockets": 300}, {"sockets": 48}], layout={"props": {second.key: {"room": 1}}}), catalog
    )
    assert next(p for p in moved.rack.plans if p.key == second.key).room == 1


def test_assign_alternates_neighbours() -> None:
    line = assign(6, 2, 0)
    assert [s for _r, _c, s in line] == [0, 1, 0, 1, 0, 1]
    grid = assign(8, 2, 4)
    by = {(r, c): s for r, c, s in grid}
    for (r, c), s in by.items():
        for nb in ((r, c + 1), (r + 1, c)):
            if nb in by:
                assert by[nb] != s
    three = {(r, c): s for r, c, s in assign(12, 3, 4)}
    for (r, c), s in three.items():
        for nb in ((r, c + 1), (r + 1, c)):
            if nb in three:
                assert three[nb] != s


def test_ap_wiring_plan(catalog: Catalog) -> None:
    site = make_site(location_code="BO115", aps=[("low_density", 40)], ap_columns=5)
    r = size_site(site, catalog)
    wifi = _items(r, "wifi_switch")
    assert len(wifi) >= 2
    w = plan_ap_wiring([("1", r)], catalog)
    room = w.rooms[0]
    assert len(room.aps) == 40 and room.shared_pairs == 0
    assert all(ap.switch >= 0 and ap.port > 0 for ap in room.aps)
    assert room.aps[0].name == "BO115-1-AP01"
    assert room.aps[0].switch != room.aps[1].switch  # left / right neighbours
    assert room.aps[0].switch != room.aps[5].switch  # the one below
    ports = [(ap.switch, ap.port) for ap in room.aps]
    assert len(set(ports)) == len(ports)  # every port used once
    assert not w.problems
    # one switch only: reported
    one = plan_ap_wiring([("1", size_site(make_site(aps=[("low_density", 6)]), catalog))], catalog)
    assert one.problems and one.rooms[0].shared_pairs > 0


def test_ap_wiring_exports(catalog: Catalog, tmp_path) -> None:
    from sitesizer.exporters.ap_wiring import ApWiringDiagram, export_ap_wiring_xlsx
    from sitesizer.exporters.diagram import export_png

    r = size_site(make_site(aps=[("low_density", 30)], ap_columns=6), catalog)
    w = plan_ap_wiring([("1", r)], catalog)
    d = ApWiringDiagram(w, title="Test")
    assert d.size().width() > 500 and d.size().height() > 200
    assert export_png(d, tmp_path / "w.png").stat().st_size > 1000
    path = export_ap_wiring_xlsx(w, tmp_path / "w.xlsx", title="Test")
    from openpyxl import load_workbook

    ws = load_workbook(path).active
    assert ws.max_row >= 30


def test_every_cabinet_has_an_optical_panel_with_an_organizer(catalog: Catalog) -> None:
    layout = {"added": ["user-1"], "props": {"user-1": {"room": 1, "size_u": 24}}}
    site = make_site(rooms=[{"sockets": 400}, {"sockets": 48}], rack_size_u=24, layout=layout)
    r = size_site(site, catalog)
    assert len(r.rack.plans) >= 4
    for p in r.rack.plans:
        items = sorted(p.items, key=lambda it: -it.u)
        odfs = [i for i, it in enumerate(items) if it.group == "fiber"]
        assert odfs, p.key
        for i in odfs:  # each ODF has its organizer right under it
            assert items[i + 1].group == "manager" and items[i + 1].u == items[i].u - 1
        assert p.used_u <= p.size_u
    fw_cab = next(p for p in r.rack.plans if any(it.group == "firewall" for it in p.items))
    assert any(it.id.endswith(":isp") for it in fw_cab.items)
