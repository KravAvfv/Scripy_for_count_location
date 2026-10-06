"""Import the company's Excel specification template ("Слаботрумка") into the catalog.

The template lists every orderable item with its internal (1C) code, name, unit and two prices
(columns A, B, C, E and G) in sections (СКС / Network / Роботи). Rows are matched to catalog
models by code, then by a model number found in the name (``Комутатор FS-148F`` → ``FS-148F``);
matched models get the code, prices, unit, section and template position, and the remaining rows
are added as new catalog items so the exported specification reproduces the template.
"""

from __future__ import annotations

import copy
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

SECTION_WORDS = (("скс", "sks"), ("network", "network"), ("мереж", "network"), ("робот", "works"), ("works", "works"))
SERVICE_WORDS = (
    "forticare",
    "fortiguard",
    "support",
    "підтримк",
    "ліценз",
    "license",
    "subscription",
    "підписк",
    "для ",
    " for ",
)
"""Rows with these words are services for a model, not the model itself."""
UNITS = {"шт", "шт.", "м", "м.", "компл", "компл.", "пос", "посл.", "послуга", "к-т", "pcs", "m"}


@dataclass
class TemplateRow:
    row: int
    code: str
    name: str
    unit: str
    price_min: float | None
    price: float | None
    section: str


@dataclass
class ImportReport:
    matched: list[tuple[str, str]] = field(default_factory=list)
    """(model, template name)"""
    added: list[str] = field(default_factory=list)
    priced: int = 0
    rows: int = 0


def _number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int | float):
        return float(value)
    text = str(value).strip().replace("\xa0", "").replace(" ", "").replace(",", ".")
    if not text or text.startswith("="):
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _section_of(text: str) -> str | None:
    low = text.lower()
    if len(low) > 40:
        return None
    for word, key in SECTION_WORDS:
        if word in low:
            return key
    return None


def read_template(path: str | Path, sheet: str | None = None) -> list[TemplateRow]:
    """Rows of the template: section headers set the section of the rows below them."""
    from openpyxl import load_workbook

    wb = load_workbook(Path(path), data_only=True, read_only=True)
    ws = None
    if sheet and sheet in wb.sheetnames:
        ws = wb[sheet]
    if ws is None:
        ws = next((wb[n] for n in wb.sheetnames if "слабо" in n.lower() or "spec" in n.lower()), wb.worksheets[0])
    rows: list[TemplateRow] = []
    section = "sks"
    for i, values in enumerate(ws.iter_rows(min_row=1, max_col=8, values_only=True), start=1):
        cells = [_text(v) for v in (*values, *[None] * (8 - len(values)))]
        filled = [c for c in cells if c]
        if not filled:
            continue
        if len(filled) == 1 and not cells[0]:
            found = _section_of(filled[0])
            if found:
                section = found
                continue
        code, name, unit = cells[0], cells[1], cells[2]
        if not name or name.lower() in ("найменування", "назва", "name"):
            continue
        if not code and unit.lower() not in UNITS:
            continue
        price_min = _number(values[4]) if len(values) > 4 else None
        price = _number(values[6]) if len(values) > 6 else None
        rows.append(TemplateRow(i, code, name, unit or "шт.", price_min, price, section))
    wb.close()
    return rows


def _match(name: str, keys: list[str]) -> str | None:
    low = name.lower()
    if any(w in low for w in SERVICE_WORDS):
        return None
    for key in keys:  # longest first: FS-148F-FPOE before FS-148F
        if re.search(rf"(?<![\w-]){re.escape(key.lower())}(?![\w-])", low):
            return key
    return None


def _new_id(row: TemplateRow, taken: set[str]) -> str:
    base = row.code or re.sub(r"\s+", " ", row.name)[:48]
    key, n = base, 2
    while key in taken:
        key = f"{base} ({n})"
        n += 1
    return key


def apply_template(data: dict[str, Any], rows: list[TemplateRow]) -> tuple[dict[str, Any], ImportReport]:
    """Return a copy of catalog ``data`` (a ``Catalog.model_dump``) updated from the template rows."""
    out = copy.deepcopy(data)
    models: dict[str, dict[str, Any]] = out["models"]
    report = ImportReport(rows=len(rows))
    by_code = {m.get("code"): k for k, m in models.items() if m.get("code")}
    keys = sorted(models, key=len, reverse=True)
    for order, row in enumerate(rows, start=1):
        key = by_code.get(row.code) if row.code else None
        key = key or _match(row.name, keys)
        if key is None:
            key = _new_id(row, set(models))
            models[key] = {
                "kind": "work" if row.section == "works" else "accessory",
                "name": {"uk": row.name, "en": row.name},
                "verified": "assumption",
                "source": "template",
            }
            report.added.append(key)
        else:
            report.matched.append((key, row.name))
        m = models[key]
        if row.code:
            m["code"] = row.code
            by_code[row.code] = key
        if row.price is not None:
            m["price"] = row.price
        if row.price_min is not None:
            m["price_min"] = row.price_min
        if row.price is not None or row.price_min is not None:
            report.priced += 1
        m["unit"] = row.unit
        m["spec_name"] = row.name
        m["section"] = row.section
        m["order"] = order
    return out, report
