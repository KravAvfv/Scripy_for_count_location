"""Passive parts for cabinets that already exist: from an Excel file or typed in.

The cabinets come from

* an Excel workbook the user already has — a rack scheme with device names like
  ``BO123-5B-ASW01 (FS-148F)`` (cabinet ``5B``, floor 5), or, when no such names are found, a
  specification with switch models and quantities (all in one cabinet);
* or a short text per cabinet: ``5B: W1 A4 V1`` (Wi-Fi, access, video switches; ``C`` core,
  ``F`` firewall).

The same house rules as the sizing engine apply: Wi-Fi switch — 1 patch panel + 1 organizer,
other switches — 2 panels + 2 organizers, an optical patch panel (+ organizer) at both ends of
every fibre link (further cabinets of a floor → the floor's first cabinet, a floor → the
firewall's cabinet), DACs per number of switches of a floor.
"""

from __future__ import annotations

import math
import re
from dataclasses import dataclass, field
from pathlib import Path

from .catalog import Catalog

ROLE_KEYS = {"W": "wifi", "A": "access", "V": "video", "C": "core", "F": "firewall"}
NAME_ROLES = {"WSW": "wifi", "ASW": "access", "VSW": "video", "CSW": "core", "FW": "firewall"}
CATEGORY_ROLES = {"wifi_switch": "wifi", "access_switch": "access", "camera_switch": "video", "core_switch": "core"}
LOOKALIKE = str.maketrans("АВСЕНКМОРТХІавсенкмортхі", "ABCEHKMOPTXIABCEHKMOPTXI")
"""Cyrillic letters that look like Latin ones (cabinet tags are often typed in Cyrillic)."""

DEVICE_RE = re.compile(r"(?<![\w-])[\w]+-(?P<tag>-?\d{1,3}[A-Z]{1,3})-(?P<role>WSW|ASW|VSW|CSW|FW)(?P<n>\d{1,3})\b")
MODEL_RE = re.compile(
    r"\b(?:FS|FG)-?\d{2,4}[A-Z]?(?:-F?POE)?\b|\bFortiSwitch[\s-]*\d{2,4}[A-Z]?(?:-F?POE)?\b", re.IGNORECASE
)
TAG_RE = re.compile(r"^(-?\d{1,3})([A-Z]{1,3})$")
SKIP_ROW = ("forticare", "support", "ліценз", "licen", "bundle", "підтримк")


@dataclass
class Cabinet:
    tag: str
    """``5B``: floor number + letter."""
    floor: int
    wifi: int = 0
    access: int = 0
    video: int = 0
    core: int = 0
    firewall: int = 0
    names: list[str] = field(default_factory=list)

    @property
    def switches(self) -> int:
        return self.wifi + self.access + self.video + self.core


@dataclass
class PassiveLine:
    model: str
    qty: int
    name: str
    why: str
    code: str = ""
    unit: str = "шт."


@dataclass
class PassiveCount:
    cabinets: list[Cabinet]
    fw_tag: str
    links: list[tuple[str, str]]
    lines: list[PassiveLine]
    source: str = ""
    warnings: list[str] = field(default_factory=list)


# =========================================================================================
# input
# =========================================================================================
def parse_tag(text: str) -> tuple[int, str] | None:
    m = TAG_RE.match(text.strip().upper().translate(LOOKALIKE))
    return (int(m.group(1)), m.group(2)) if m else None


def parse_cabinet(text: str) -> Cabinet:
    """``5B: W1 A4 V1`` (or ``5B:W1,A4,V1``) → a cabinet. Raises ``ValueError`` on bad input."""
    if ":" not in text:
        raise ValueError(f"Очікується «5B: W1 A4 V1», отримано «{text}»")
    head, body = text.split(":", 1)
    tag = parse_tag(head)
    if tag is None:
        raise ValueError(f"Тег шафи «{head.strip()}» — очікується номер поверху й літера, напр. 5B")
    cab = Cabinet(tag=f"{tag[0]}{tag[1]}", floor=tag[0])
    for token in re.split(r"[\s,;]+", body.strip().upper()):
        if not token:
            continue
        m = re.fullmatch(r"([WAVCF])(\d+)", token)
        if not m:
            raise ValueError(f"Незрозуміло «{token}»: W — Wi-Fi, A — доступ, V — відео, C — ядро, F — фаєрвол")
        role = ROLE_KEYS[m.group(1)]
        setattr(cab, role, getattr(cab, role) + int(m.group(2)))
    return cab


def _normalize_model(text: str) -> str:
    t = text.upper().replace(" ", "")
    t = re.sub(r"^FORTISWITCH-?", "FS-", t)
    return re.sub(r"^(FS|FG)(\d)", r"\1-\2", t)


def cabinets_from_xlsx(path: str | Path, catalog: Catalog) -> tuple[list[Cabinet], list[str], str]:
    """Read cabinets from a workbook: (cabinets, warnings, how they were found)."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    names: dict[str, tuple[str, str]] = {}
    rows: list[list[object]] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = list(row)
            rows.append(cells)
            for v in cells:
                if not isinstance(v, str):
                    continue
                text = v.upper().translate(LOOKALIKE)
                for m in DEVICE_RE.finditer(text):
                    names[m.group(0)] = (m.group("tag"), NAME_ROLES[m.group("role")])
    wb.close()
    warnings: list[str] = []
    if names:
        cabs: dict[str, Cabinet] = {}
        for full, (tag, role) in sorted(names.items()):
            parsed = parse_tag(tag)
            if parsed is None:
                continue
            cab = cabs.setdefault(tag, Cabinet(tag=tag, floor=parsed[0]))
            setattr(cab, role, getattr(cab, role) + 1)
            cab.names.append(full)
        return _sorted(list(cabs.values())), warnings, "names"

    # no device names: a specification with models and quantities
    role_of: dict[str, str] = {}
    for key, role_name in CATEGORY_ROLES.items():
        if key in catalog.categories:
            cat = catalog.categories[key]
            role_of.setdefault(cat.base, role_name)
            role_of.setdefault(cat.premium, role_name)
    cab = Cabinet(tag="1A", floor=1)
    found = False
    for cells in rows:
        texts = [c for c in cells if isinstance(c, str)]
        line = " ".join(texts)
        if not line or any(s in line.lower() for s in SKIP_ROW):
            continue
        hit = MODEL_RE.search(line.upper())
        if not hit:
            continue
        qty = _row_qty(cells)
        if not qty:
            continue
        model = _normalize_model(hit.group(0))
        if model.startswith("FG-"):
            cab.firewall += qty
        else:
            kind = role_of.get(model) or ("core" if model.startswith("FS-10") else "")
            if not kind:
                dev = catalog.models.get(model)
                ports = dev.ports.count if dev and dev.ports else 48
                kind = "wifi" if "POE" in model and ports <= 24 else ("video" if "POE" in model else "access")
            setattr(cab, kind, getattr(cab, kind) + qty)
        cab.names.append(f"{model} × {qty}")
        found = True
    if found:
        warnings.append(
            "У файлі немає назв пристроїв на кшталт BO123-5B-ASW01 — усі комутатори з кількостей "
            "специфікації пораховано в одній шафі (1A). Для поверхів і шаф дайте схему шаф."
        )
        return [cab], warnings, "spec"
    return [], ["У файлі не знайдено ні назв пристроїв (BO123-5B-ASW01), ні моделей FS-/FG- з кількістю."], ""


def _row_qty(cells: list[object]) -> int:
    """Quantity of a specification row: column D of the company template, else the first whole number."""
    if len(cells) > 3 and isinstance(cells[3], int | float) and not isinstance(cells[3], bool) and cells[3] > 0:
        return int(cells[3])
    for c in cells:
        if isinstance(c, int | float) and not isinstance(c, bool) and 0 < c < 1000 and float(c).is_integer():
            return int(c)
    return 0


def _sorted(cabs: list[Cabinet]) -> list[Cabinet]:
    return sorted(cabs, key=lambda c: (-c.floor, len(c.tag), c.tag))


# =========================================================================================
# counting
# =========================================================================================
def fibre_links(cabinets: list[Cabinet], fw_tag: str) -> list[tuple[str, str]]:
    """(child, parent): further cabinets of a floor → the floor's first one; a floor → the firewall."""
    floors: dict[int, list[Cabinet]] = {}
    for c in cabinets:
        floors.setdefault(c.floor, []).append(c)
    fw = next((c for c in cabinets if c.tag == fw_tag), cabinets[0] if cabinets else None)
    links: list[tuple[str, str]] = []
    for n, cabs in floors.items():
        head = fw if fw is not None and fw.floor == n else sorted(cabs, key=lambda c: (len(c.tag), c.tag))[0]
        for c in cabs:
            if c is not head:
                links.append((c.tag, head.tag))
        if fw is not None and head is not fw:
            links.append((head.tag, fw.tag))
    return links


def count_passive(
    cabinets: list[Cabinet],
    catalog: Catalog,
    fw_tag: str = "",
    sockets: int = 0,
    cameras: int = 0,
    aps: int = 0,
    avg_run_m: int = 0,
) -> PassiveCount:
    """Passive parts of the given cabinets (and the copper, when the endpoint counts are known)."""
    pr, rules = catalog.passive, catalog.rules
    fw_tag = (
        fw_tag or next((c.tag for c in cabinets if c.firewall), "") or (cabinets[0].tag if cabinets else "")
    ).upper()
    fw_tag = fw_tag.translate(LOOKALIKE)
    links = fibre_links(cabinets, fw_tag)
    per = pr.panels_per_switch

    def orgs(panels: int) -> int:
        return 1 + (1 if panels - math.ceil(panels / 2) > 0 else 0)

    pw, pa, pv = per.get("wifi_switch", 1), per.get("access_switch", 2), per.get("camera_switch", 2)
    panels = sum(c.wifi * pw + c.access * pa + c.video * pv for c in cabinets)
    odf = 2 * len(links)
    managers = sum(c.wifi * orgs(pw) + c.access * orgs(pa) + c.video * orgs(pv) for c in cabinets) + odf

    floors: dict[int, int] = {}
    for c in cabinets:
        floors[c.floor] = floors.get(c.floor, 0) + c.switches
    dac_short = sum(math.ceil(n / rules.dac_short_per_switches) for n in floors.values() if n)
    dac_long = sum(math.ceil(n / rules.dac_long_per_switches) for n in floors.values() if n)

    by_tag = {c.tag: c for c in cabinets}
    lengths = []
    for a, b in links:
        fa, fb = by_tag[a].floor, by_tag[b].floor
        lengths.append(pr.fiber_cabinet_m if fa == fb else max(1, abs(fa - fb)) * pr.floor_height_m)
    longest = max(lengths, default=0) + pr.fiber_slack_m
    fiber = None
    for _key, spec in sorted(pr.fiber.items(), key=lambda kv: kv[1].max_10g_m):
        if longest <= spec.max_10g_m:
            fiber = spec
            break
    if fiber is None and pr.fiber:
        fiber = max(pr.fiber.values(), key=lambda s: s.max_10g_m)

    lines: list[PassiveLine] = []

    def add(model: str, qty: int, why: str) -> None:
        if qty <= 0 or not model:
            return
        dev = catalog.models.get(model)
        name = (dev.name.get("uk") or next(iter(dev.name.values()), "")) if dev else ""
        lines.append(PassiveLine(model, qty, name, why, dev.code if dev else "", dev.unit if dev else "шт."))

    ports = panels * pr.panel_ports
    add(pr.panel, panels, f"Wi-Fi комутатор — {pw} ПП, інші — по {pa}")
    add(pr.manager, managers, f"Wi-Fi комутатор — {orgs(pw)}, інші — по {orgs(pa)}, оптична ПП — 1")
    if sockets or cameras or aps:
        links_cu = sockets + cameras + aps
        avg = avg_run_m or rules.avg_cable_run_m_default
        cable_m = links_cu * (avg + pr.cable_slack_m)
        add(
            pr.cable,
            math.ceil(cable_m / pr.cable_drum_m),
            f"{links_cu} ліній × ({avg} + {pr.cable_slack_m}) м = {cable_m} м",
        )
        add(pr.jack, math.ceil(links_cu * 2 / pr.jack_pack), f"{links_cu} ліній × 2 модулі, упаковки по {pr.jack_pack}")
        add(pr.cord_rack, links_cu + cameras + aps, "панель ↔ комутатор + біля AP/камер")
        add(pr.cord_user, sockets, "від розетки до пристрою")
    else:
        add(pr.jack, math.ceil(ports / pr.jack_pack), f"модулі в усі порти ПП: {ports}, упаковки по {pr.jack_pack}")
        add(pr.cord_rack, ports, f"патч-корди панель ↔ комутатор: {ports} портів ПП")
    if fiber is not None and links:
        add(fiber.housing_12, odf, f"{len(links)} оптичних лінків × 2 кінці")
        add(fiber.transceiver, odf, "трансивер на кожен кінець лінку")
        add(fiber.cord, odf, "оптичний патч-корд на кожну панель")
        add(pr.splice_protector, odf * fiber.cable_fibers, f"{fiber.cable_fibers} волокон на кожну панель")
        add(fiber.cable, sum(lengths) + pr.fiber_slack_m * len(links), "траси між шафами / до фаєрвола + запас")
    add(rules.dac_short_model, dac_short, f"по 1 на кожні {rules.dac_short_per_switches} комутатори поверху")
    add(rules.dac_long_model, dac_long, f"по 1 на кожні {rules.dac_long_per_switches} комутаторів поверху")
    return PassiveCount(cabinets=cabinets, fw_tag=fw_tag, links=links, lines=lines)


# =========================================================================================
# output
# =========================================================================================
def format_count(pc: PassiveCount) -> str:
    out = ["Шафи:"]
    for c in pc.cabinets:
        parts = [f"Wi-Fi {c.wifi}", f"доступ {c.access}", f"відео {c.video}"]
        if c.core:
            parts.append(f"ядро {c.core}")
        fw = " · фаєрвол" if c.tag == pc.fw_tag else ""
        out.append(f"  {c.tag}: " + ", ".join(parts) + fw)
    if pc.links:
        out.append("Оптичні лінки: " + ", ".join(f"{a}→{b}" for a, b in pc.links))
    out.append("")
    width = max((len(line.model) for line in pc.lines), default=10)
    for line in pc.lines:
        out.append(f"  {line.model:<{width}}  {line.qty:>6}  {line.name}")
        out.append(f"  {'':<{width}}          ↳ {line.why}")
    for w in pc.warnings:
        out.append(f"! {w}")
    return "\n".join(out)


def write_xlsx(pc: PassiveCount, path: str | Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Пасивка"
    ws.append(["Код", "Найменування", "Модель", "Од.", "К-сть", "Як пораховано"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color="1F4E78", end_color="1F4E78")
    for line in pc.lines:
        ws.append([line.code or None, line.name, line.model, line.unit, line.qty, line.why])
    ws.append([])
    ws.append(["", "Шафи"])
    for cab in pc.cabinets:
        fw = " (фаєрвол)" if cab.tag == pc.fw_tag else ""
        ws.append(["", f"{cab.tag}{fw}: Wi-Fi {cab.wifi}, доступ {cab.access}, відео {cab.video}, ядро {cab.core}"])
    if pc.links:
        ws.append(["", "Оптичні лінки: " + ", ".join(f"{a}→{b}" for a, b in pc.links)])
    for w in pc.warnings:
        ws.append(["", w])
    for col, width in zip("ABCDEF", (12, 70, 22, 6, 8, 60), strict=True):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=c.column in (2, 6))
    path = Path(path)
    wb.save(path)
    return path
