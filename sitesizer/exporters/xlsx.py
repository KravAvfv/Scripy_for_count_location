"""Excel export in the company's format.

Sheets (the user picks which ones):

* **Слаботрумка** — the specification exactly like the Excel template: code (1C), name, unit,
  quantity, two prices with ``=D*E`` / ``=D*G`` formulas, grey section rows (СКС / Network / Роботи)
  and every template item listed (0 when unused);
* **<site> - Схема+шафи** — the cabinets drawn in cells (unit numbers on both sides, switches
  green, organizers grey) plus the topology picture;
* **<site> - IP** — VLAN table;
* **Ціни та кількість** — the bill of materials with quantities and prices (manual edits included);
* **Вихідні дані** — inputs and checks.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .. import __version__
from ..core.catalog import Catalog
from ..core.models import RackPlan, SiteResult
from ..core.pricing import summarize_prices
from ..core.report import BOM_GROUP_ORDER, ip_columns, ip_rows, spec_sections
from ..i18n import Translator

NAVY = "1F4E78"
LIGHT = "EAF0F6"
BORDER = "B7B7B7"
MUTED = "7C8894"
TEXT = "1F2D3A"
FONT = "Arial"

THIN = Side(style="thin", color=BORDER)
BLACK = Side(style="thin", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOX_DARK = Border(left=BLACK, right=BLACK, top=BLACK, bottom=BLACK)
HEADER_FILL = PatternFill("solid", start_color=NAVY, end_color=NAVY)
GROUP_FILL = PatternFill("solid", start_color=LIGHT, end_color=LIGHT)
SECTION_FILL = PatternFill("solid", start_color="BFBFBF", end_color="BFBFBF")
WRAP_TOP = Alignment(vertical="top", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
MONEY = "#,##0.00"

SHEETS = ("spec", "racks", "ip", "prices", "inputs")
"""Everything :func:`export_xlsx` can write; ``options["sheets"]`` picks a subset."""

RACK_FILLS = {
    "wifi_switch": "33CC33",
    "access_switch": "33CC33",
    "camera_switch": "33CC33",
    "core_switch": "33CC33",
    "firewall": "F4B183",
    "manager": "A6A6A6",
    "power": "FFE699",
    "pdu": "FFF2CC",
    "fiber": "DDEBF7",
    "shelf": "D9D9D9",
    "blank": "EDEDED",
    "device": "9DC3E6",
}


class _PngImage(XLImage):
    """openpyxl image from PNG bytes without requiring Pillow."""

    def __init__(self, data: bytes, width: int, height: int) -> None:
        self.ref = None
        self._bytes = data
        self.width = width
        self.height = height
        self.format = "png"

    def _data(self) -> bytes:
        return self._bytes


def sheet_title(text: str) -> str:
    """Excel sheet names: ≤ 31 characters, no ``[]:*?/\\``."""
    cleaned = "".join("_" if c in "[]:*?/\\" else c for c in text).strip() or "Sheet"
    return cleaned[:31]


def _title(ws: Worksheet, text: str, subtitle: str, ncols: int) -> int:
    ws.cell(row=1, column=1, value=text).font = Font(name=FONT, size=14, bold=True, color=TEXT)
    ws.cell(row=2, column=1, value=subtitle).font = Font(name=FONT, size=9, color=MUTED)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    ws.row_dimensions[1].height = 22
    return 4


def _header(ws: Worksheet, row: int, columns: list[str]) -> None:
    for c, name in enumerate(columns, start=1):
        cell = ws.cell(row=row, column=c, value=name)
        cell.fill = HEADER_FILL
        cell.font = Font(name=FONT, color="FFFFFF", bold=True, size=10)
        cell.alignment = CENTER
        cell.border = BOX
    ws.row_dimensions[row].height = 30


def _autofit(ws: Worksheet, first_row: int, ncols: int, caps: dict[int, int] | None = None) -> None:
    caps = caps or {}
    widths = [10] * ncols
    for row in ws.iter_rows(min_row=first_row, max_col=ncols):
        for cell in row:
            if cell.value is None or type(cell).__name__ == "MergedCell":
                continue
            longest = max((len(part) for part in str(cell.value).split("\n")), default=0)
            widths[cell.column - 1] = max(widths[cell.column - 1], longest + 2)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(w, caps.get(i, 90))


def _print_setup(ws: Worksheet, header_row: int | None, landscape: bool = True) -> None:
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    if header_row:
        ws.print_title_rows = f"{header_row}:{header_row}"
    ws.page_margins.left = ws.page_margins.right = 0.4


def _meta_line(opts: dict[str, Any], result: SiteResult) -> str:
    s = result.input
    items = [
        opts.get("project"),
        opts.get("customer"),
        s.location_code and f"{s.location_code}" + (f" · ID {s.location_id}" if s.location_id is not None else ""),
        date.today().strftime("%d.%m.%Y"),
        opts.get("author"),
        opts.get("company"),
    ]
    return "  ·  ".join(str(m) for m in items if m)


def export_xlsx(
    result: SiteResult,
    path: str | Path,
    catalog: Catalog | None = None,
    lang: str = "uk",
    options: dict[str, Any] | None = None,
) -> Path:
    """Write the workbook and return its path.

    ``options``: ``sheets`` (subset of :data:`SHEETS`), ``only_used`` (drop zero rows from the
    specification), ``project``, ``customer``, ``author``, ``company``, ``discount_pct``, ``vat_pct``.
    """
    from ..core.catalog import load_default_catalog

    catalog = catalog or load_default_catalog()
    t = Translator(lang)
    opts = options or {}
    sheets = [s for s in SHEETS if s in set(opts.get("sheets") or SHEETS)] or ["spec"]
    path = Path(path)
    wb = Workbook()
    default = wb.active
    site = result.input.location_code or result.input.name

    for key in sheets:
        if key == "spec":
            _spec_sheet(wb.create_sheet(sheet_title(t.t("xl.sheet_spec"))), result, catalog, t, opts)
        elif key == "racks" and (result.rack.plans or result.has_equipment):
            _racks_sheet(wb.create_sheet(sheet_title(t.t("xl.sheet_scheme", site=site))), result, catalog, t, lang)
        elif key == "ip" and result.ip_plan is not None and result.ip_plan.segments:
            _ip_sheet(wb.create_sheet(sheet_title(t.t("xl.sheet_ip_site", site=site))), result, t, opts)
        elif key == "prices":
            _prices_sheet(wb.create_sheet(sheet_title(t.t("xl.sheet_prices"))), result, catalog, t, opts)
        elif key == "inputs":
            _inputs_sheet(wb.create_sheet(sheet_title(t.t("xl.sheet_inputs"))), result, catalog, t)
    if default is not None and len(wb.sheetnames) > 1:
        wb.remove(default)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# =========================================================================================
# Слаботрумка — the specification template
# =========================================================================================
def _spec_sheet(ws: Worksheet, result: SiteResult, catalog: Catalog, t: Translator, opts: dict[str, Any]) -> None:
    cols = [
        t.t("xl.col_code"),
        t.t("xl.col_name"),
        t.t("xl.col_unit"),
        t.t("xl.col_qty"),
        t.t("xl.col_price_min"),
        t.t("xl.col_sum_min"),
        t.t("xl.col_price"),
        t.t("xl.col_sum"),
    ]
    row = _title(ws, t.t("xl.spec_title", site=result.input.name), _meta_line(opts, result), len(cols))
    header_row = row
    _header(ws, row, cols)
    row += 1
    first = row
    for section, rows in spec_sections(result, catalog, t, bool(opts.get("only_used"))):
        cell = ws.cell(row=row, column=1, value=t.t(f"section.{section}"))
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(cols))
        cell.font = Font(name=FONT, size=10, bold=True, italic=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        for c in range(1, len(cols) + 1):
            ws.cell(row=row, column=c).fill = SECTION_FILL
            ws.cell(row=row, column=c).border = BOX
        row += 1
        for sr in rows:
            values: list[Any] = [
                sr.code or None,
                sr.name,
                sr.unit,
                sr.qty,
                sr.price_min,
                f"=D{row}*E{row}",
                sr.price,
                f"=D{row}*G{row}",
            ]
            for c, v in enumerate(values, start=1):
                cell = ws.cell(row=row, column=c, value=v)
                cell.border = BOX
                cell.font = Font(name=FONT, size=10, bold=bool(sr.qty) and c in (2, 4))
                if c in (3, 4):
                    cell.alignment = Alignment(horizontal="center", vertical="top")
                elif c == 2:
                    cell.alignment = WRAP_TOP
                if c >= 5:
                    cell.number_format = MONEY
                if c == 1:
                    cell.number_format = "@"
            row += 1
    last = row - 1
    row += 1
    totals = [
        (t.t("xl.total"), f"=SUM(F{first}:F{last})", f"=SUM(H{first}:H{last})"),
    ]
    vat = float(opts.get("vat_pct", 0) or 0)
    if vat:
        totals.append((t.t("price.vat", pct=f"{vat:g}"), f"=F{row}*{vat / 100:g}", f"=H{row}*{vat / 100:g}"))
        totals.append((t.t("xl.total_vat"), f"=F{row}+F{row + 1}", f"=H{row}+H{row + 1}"))
    for i, (label, f_sum, h_sum) in enumerate(totals):
        strong = i == len(totals) - 1
        ws.cell(row=row, column=2, value=label).font = Font(name=FONT, bold=strong, size=10)
        ws.cell(row=row, column=2).alignment = Alignment(horizontal="right")
        for c, v in ((6, f_sum), (8, h_sum)):
            cell = ws.cell(row=row, column=c, value=v)
            cell.number_format = MONEY
            cell.font = Font(name=FONT, bold=strong, size=10)
            cell.border = BOX
        row += 1
    for i, w in enumerate((12, 72, 7, 8, 13, 15, 13, 15), start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"C{first}"
    ws.auto_filter.ref = f"A{header_row}:H{last}"
    _print_setup(ws, header_row, landscape=False)


# =========================================================================================
# Схема + шафи — cabinets drawn in cells, plus the topology picture
# =========================================================================================
def rack_cell_label(item_label: str, model: str, group: str) -> str:
    if group in ("wifi_switch", "access_switch", "camera_switch", "core_switch", "firewall", "device") and model:
        return f"{item_label} ({model})" if model not in item_label else item_label
    return item_label


def _racks_sheet(ws: Worksheet, result: SiteResult, catalog: Catalog, t: Translator, lang: str) -> None:
    ws.sheet_view.showGridLines = False
    ws.cell(row=1, column=2, value=t.t("xl.racks_title", site=result.input.name)).font = Font(
        name=FONT, size=14, bold=True, color=TEXT
    )
    top = 3
    col = 2
    max_rows = 0
    for plan in result.rack.plans:
        _rack_block(ws, plan, top, col)
        max_rows = max(max_rows, plan.size_u + 2)
        col += 4
    row = top + max_rows + 2
    if result.has_equipment:
        try:
            png, w, h = _diagram_png(result, catalog, lang)
        except Exception:
            png = None
        if png:
            ws.cell(row=row, column=2, value=t.t("xl.sheet_diagram")).font = Font(name=FONT, size=12, bold=True)
            scale = min(1.0, 1000 / w)
            ws.add_image(_PngImage(png, int(w * scale), int(h * scale)), f"B{row + 1}")
    _print_setup(ws, None, landscape=False)


def _rack_block(ws: Worksheet, plan: RackPlan, top: int, col: int) -> None:
    """One cabinet: title, black top bar, unit numbers on both sides, items in the middle column."""
    left, mid, right = col, col + 1, col + 2
    ws.column_dimensions[get_column_letter(left)].width = 5
    ws.column_dimensions[get_column_letter(mid)].width = 44
    ws.column_dimensions[get_column_letter(right)].width = 5
    ws.column_dimensions[get_column_letter(col + 3)].width = 3
    title = ws.cell(row=top, column=left, value=f"{plan.name} · {plan.size_u}U")
    ws.merge_cells(start_row=top, start_column=left, end_row=top, end_column=right)
    title.font = Font(name=FONT, size=11, bold=True)
    title.alignment = CENTER
    ws.row_dimensions[top].height = 30
    for c in (left, mid, right):
        ws.cell(row=top, column=c).border = BOX_DARK
        ws.cell(row=top + 1, column=c).fill = PatternFill("solid", start_color="000000", end_color="000000")
    first = top + 2
    by_top = {it.top: it for it in plan.items if 1 <= it.u <= plan.size_u}
    for i in range(plan.size_u):
        u = plan.size_u - i
        r = first + i
        for c in (left, right):
            cell = ws.cell(row=r, column=c, value=u)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = CENTER
            cell.border = BOX_DARK
        ws.cell(row=r, column=mid).border = BOX_DARK
        ws.row_dimensions[r].height = 18
        it = by_top.get(u)
        if it is None:
            continue
        cell = ws.cell(row=r, column=mid, value=rack_cell_label(it.label, it.model, it.group))
        cell.font = Font(name=FONT, size=10, bold=it.group in RACK_FILLS and it.group != "manager")
        cell.alignment = CENTER
        fill = RACK_FILLS.get(it.group)
        span = min(it.height, plan.size_u - i)
        if span > 1:
            ws.merge_cells(start_row=r, start_column=mid, end_row=r + span - 1, end_column=mid)
        if fill:
            for k in range(span):
                ws.cell(row=r + k, column=mid).fill = PatternFill("solid", start_color=fill, end_color=fill)
                ws.cell(row=r + k, column=mid).border = BOX_DARK


# =========================================================================================
# IP
# =========================================================================================
def _ip_sheet(ws: Worksheet, result: SiteResult, t: Translator, opts: dict[str, Any]) -> None:
    plan = result.ip_plan
    assert plan is not None
    with_addr = plan.has_addresses
    cols = ip_columns(t, with_addr)
    sub = plan.base_network or t.t("xl.ip_no_base")
    if plan.location_id is not None:
        sub = t.t("xl.ip_by_id", id=plan.location_id, template=plan.template)
    r = _title(ws, t.t("xl.ip_title", site=result.input.name), sub, len(cols))
    _header(ws, r, cols)
    hdr = r
    r += 1
    for values in ip_rows(result):
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BOX
            cell.alignment = WRAP_TOP
            cell.font = Font(name=FONT, size=10)
        r += 1
    if plan.error:
        ws.cell(row=r + 1, column=1, value=plan.error).font = Font(name=FONT, color="A3433B", bold=True)
    _autofit(ws, hdr, len(cols), {len(cols): 60})
    ws.freeze_panes = ws.cell(row=hdr + 1, column=1)
    _print_setup(ws, hdr)


# =========================================================================================
# prices & quantities
# =========================================================================================
def _prices_sheet(ws: Worksheet, result: SiteResult, catalog: Catalog, t: Translator, opts: dict[str, Any]) -> None:
    cols = [
        t.t("col.category"),
        t.t("col.model"),
        t.t("xl.col_code"),
        t.t("col.description"),
        t.t("xl.col_unit"),
        t.t("xl.col_qty"),
        t.t("col.unit_price"),
        t.t("col.total_price"),
    ]
    row = _title(ws, t.t("xl.prices_title", site=result.input.name), _meta_line(opts, result), len(cols))
    header_row = row
    _header(ws, row, cols)
    row += 1
    first = row
    order = {g: i for i, g in enumerate(BOM_GROUP_ORDER)}
    grouped: dict[str, list[Any]] = {}
    for line in result.bom:
        if line.qty is None:
            continue
        grouped.setdefault(line.group, []).append(line)
    for g in sorted(grouped, key=lambda g: order.get(g, 99)):
        cell = ws.cell(row=row, column=1, value=t.t(f"group.{g}"))
        cell.font = Font(name=FONT, bold=True, color=NAVY)
        for c in range(1, len(cols) + 1):
            ws.cell(row=row, column=c).fill = GROUP_FILL
            ws.cell(row=row, column=c).border = BOX
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(cols))
        row += 1
        for line in grouped[g]:
            values: list[Any] = [
                line.category,
                line.model,
                line.code or None,
                line.description,
                line.unit,
                line.qty,
                line.unit_price,
                f"=F{row}*G{row}",
            ]
            for c, v in enumerate(values, start=1):
                cell = ws.cell(row=row, column=c, value=v)
                cell.border = BOX
                cell.alignment = CENTER if c in (5, 6) else WRAP_TOP
                cell.font = Font(name=FONT, size=10, bold=c == 2, italic=line.manual and c == 6)
                if c >= 7:
                    cell.number_format = MONEY
            row += 1
    last = row - 1
    row += 1
    s = summarize_prices(
        result.bom,
        catalog.meta.currency,
        float(opts.get("discount_pct", 0) or 0),
        float(opts.get("vat_pct", 0) or 0),
    )
    items = [(t.t("price.subtotal"), f"=SUM(H{first}:H{last})")]
    if s and s.discount_pct:
        items.append((t.t("price.discount", pct=f"{s.discount_pct:g}"), f"=-H{row}*{s.discount_pct / 100:g}"))
    if s and s.vat_pct:
        base = f"(H{row}+H{row + 1})" if s.discount_pct else f"H{row}"
        items.append((t.t("price.vat", pct=f"{s.vat_pct:g}"), f"={base}*{s.vat_pct / 100:g}"))
    if len(items) > 1:
        items.append((t.t("price.total"), f"=SUM(H{row}:H{row + len(items) - 1})"))
    for i, (k, v) in enumerate(items):
        strong = i == len(items) - 1
        ws.cell(row=row, column=7, value=k).font = Font(name=FONT, bold=strong)
        c = ws.cell(row=row, column=8, value=v)
        c.number_format = MONEY
        c.font = Font(name=FONT, bold=strong)
        row += 1
    if s and not s.complete:
        ws.cell(row=row, column=1, value=t.t("price.incomplete", n=s.unpriced_lines)).font = Font(
            name=FONT, size=9, italic=True, color=MUTED
        )
    for i, w in enumerate((26, 20, 12, 50, 7, 8, 14, 16), start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{first}"
    _print_setup(ws, header_row)


# =========================================================================================
# inputs
# =========================================================================================
def _inputs_sheet(ws: Worksheet, result: SiteResult, catalog: Catalog, t: Translator) -> None:
    r = _title(ws, t.t("xl.inputs_title"), result.input.name, 2)
    _header(ws, r, [t.t("xl.param"), t.t("xl.value")])
    hdr = r
    r += 1
    for k, v in _input_rows(result, t, catalog):
        ws.cell(row=r, column=1, value=k).border = BOX
        c = ws.cell(row=r, column=2, value=v)
        c.border = BOX
        c.alignment = WRAP_TOP
        r += 1
    r += 1
    ws.cell(row=r, column=1, value=t.t("xl.checks")).font = Font(name=FONT, bold=True, color=NAVY)
    r += 1
    for chk in result.checks:
        ws.cell(row=r, column=1, value=t.t(f"xl.sev.{chk.severity.value}"))
        c = ws.cell(row=r, column=2, value=chk.message + (f"\n{chk.hint}" if chk.hint else ""))
        c.alignment = WRAP_TOP
        r += 1
    r += 1
    note = ws.cell(row=r, column=1, value=t.t("xl.footer", v=__version__, updated=catalog.meta.updated))
    note.font = Font(name=FONT, size=9, italic=True, color=MUTED)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 100
    ws.freeze_panes = ws.cell(row=hdr + 1, column=1)


def _diagram_png(result: SiteResult, catalog: Catalog, lang: str) -> tuple[bytes, int, int]:
    from PySide6.QtCore import QBuffer, QByteArray, QIODevice

    from .diagram import SiteDiagram, render_image

    img = render_image(SiteDiagram(result, catalog, lang, title=False), scale=1.6)
    data = QByteArray()
    buf = QBuffer(data)
    buf.open(QIODevice.OpenModeFlag.WriteOnly)
    img.save(buf, "PNG")  # type: ignore[call-overload]
    buf.close()
    ratio = img.devicePixelRatio()
    return bytes(data.data()), int(img.width() / ratio), int(img.height() / ratio)


def _input_rows(result: SiteResult, t: Translator, catalog: Catalog | None) -> list[tuple[str, Any]]:
    s = result.input
    tier_label = t.pick(catalog.tier(s.tier).label) if catalog else str(s.tier)
    rows: list[tuple[str, Any]] = [
        (t.t("ui.site_name"), s.name),
        (t.t("ui.location_code"), s.location_code or "—"),
        (t.t("ui.location_id"), s.location_id if s.location_id is not None else "—"),
        (t.t("ui.floors"), s.floors),
        (t.t("ui.mode"), t.t("ui.mode_quick") if s.mode == "quick" else t.t("ui.mode_extended")),
        (t.t("ui.sockets"), s.sockets),
        (t.t("ui.cameras"), s.cameras),
    ]
    for g in s.ap_groups:
        zone = catalog.ap_zones.get(g.zone) if catalog else None
        name = g.name or (t.pick(zone.label) if zone else g.zone)
        rows.append((t.t("xl.zone", name=name), f"{g.qty} × {g.model or (zone.model if zone else '')}"))
    rows += [
        (t.t("ui.tier"), f"{s.tier} · {tier_label}"),
        (t.t("ui.psu"), t.t("xl.yes") if result.dual_psu else t.t("xl.no")),
        (t.t("ui.agg"), t.t(f"ui.agg_{s.aggregation}")),
        (t.t("ui.reserve"), f"+{round((result.counts.reserve_factor - 1) * 100)}%" if s.reserve else t.t("xl.no")),
        (t.t("ui.fortios"), result.fortios_version),
        (t.t("ui.rack_size"), f"{s.rack_size_u}U" if s.rack_size_u else t.t("ui.rack_size_auto")),
    ]
    if s.mode == "extended":
        rows += [
            (t.t("ui.clients"), s.wifi_clients_expected or t.t("xl.auto")),
            (t.t("ui.guest"), s.guest_clients),
            (t.t("ui.iot"), s.iot_devices),
            (t.t("ui.cam_w"), s.camera_watts or (catalog.rules.camera_watts_default if catalog else "")),
            (t.t("ui.inspected"), s.inspected_mbps or "—"),
            (t.t("ui.max_run"), s.max_cable_run_m or "—"),
            (t.t("ui.fiber_type"), t.t("ui.fiber_auto") if s.fiber_type == "auto" else s.fiber_type.upper()),
            (t.t("ui.fiber_backbone"), s.fiber_backbone_m or t.t("xl.auto")),
            (t.t("ui.base_net"), s.ip.base_network or "—"),
        ]
    rows += [
        (t.t("ui.pw_total"), f"{result.power.total_w:.0f} W · UPS ≥ {result.power.ups_va} VA"),
    ]
    if s.notes:
        rows.append((t.t("ui.notes"), s.notes))
    return rows
