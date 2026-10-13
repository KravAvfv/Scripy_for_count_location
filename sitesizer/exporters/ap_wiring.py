"""Painter and Excel table of the AP wiring scheme (:mod:`sitesizer.core.ap_wiring`).

Each telecom room: the access points as they stand on the plan (row by row), coloured by the
switch they go to, then a card per Wi-Fi switch with its ports. Same ``size`` / ``paint`` interface
as the other diagrams, so the app's canvas, PNG / SVG and PDF use it as is.
"""

from __future__ import annotations

from pathlib import Path

from PySide6.QtCore import QRectF, QSizeF, Qt
from PySide6.QtGui import QColor, QFont, QFontMetricsF, QPainter, QPen

from ..core.ap_wiring import ApWiring, WiringRoom
from ..i18n import Translator
from .diagram import EXPORT_STYLE, DiagramStyle, _font, ensure_gui_app

SWITCH_COLORS = ("#2F6DB5", "#D9822B", "#2E8B57", "#8E5BB5", "#C0392B", "#16A2A2", "#8D6E63", "#C2557A")
PAD = 32.0
CELL_W = 122.0
CELL_H = 50.0
GAP = 8.0
CARD_W = 236.0
LINE_H = 16.0
WRAP = 10
"""APs per row when the plan has no rows (one line): only for the drawing."""


def switch_color(i: int) -> str:
    return SWITCH_COLORS[i % len(SWITCH_COLORS)] if i >= 0 else "#9AA7B4"


def short_name(label: str) -> str:
    """``BO115-1A-WSW01`` -> ``1A-WSW01`` (the site code is the same everywhere)."""
    parts = label.split("-")
    return "-".join(parts[-2:]) if len(parts) > 2 else label


class ApWiringDiagram:
    def __init__(self, wiring: ApWiring, lang: str = "uk", style: DiagramStyle = EXPORT_STYLE, title: str = "") -> None:
        ensure_gui_app()
        self.w = wiring
        self.t = Translator(lang)
        self.s = style
        self.title = title
        self.f_title = _font(18, QFont.Weight.DemiBold)
        self.f_room = _font(14, QFont.Weight.DemiBold)
        self.f_name = _font(12, QFont.Weight.DemiBold)
        self.f_small = _font(10.5)
        self.f_caption = _font(11.5, italic=True)
        self._size = self._walk(None)

    def size(self) -> QSizeF:
        return self._size

    def paint(self, p: QPainter, background: bool = True) -> None:
        p.setRenderHint(QPainter.RenderHint.Antialiasing)
        p.setRenderHint(QPainter.RenderHint.TextAntialiasing)
        if background:
            p.fillRect(QRectF(0, 0, self._size.width(), self._size.height()), QColor(self.s.bg))
        self._walk(p)

    # ---- one pass for both measuring and drawing ----------------------------------------------
    def _walk(self, p: QPainter | None) -> QSizeF:
        s = self.s
        y = PAD
        width = 520.0
        if self.title:
            if p:
                p.setPen(QColor(s.light_text))
                p.setFont(self.f_title)
                p.drawText(
                    QRectF(PAD, y, 2000, 28), Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter, self.title
                )
            y += 38
        if p:
            p.setPen(QColor(s.muted))
            p.setFont(self.f_caption)
            p.drawText(QRectF(PAD, y, 2000, 18), Qt.AlignmentFlag.AlignLeft, self.t.t("apw.rule"))
        y += 30
        width = max(
            width,
            PAD + QFontMetricsF(self.f_caption).horizontalAdvance(self.t.t("apw.rule")),
            PAD + QFontMetricsF(self.f_title).horizontalAdvance(self.title),
            *(PAD + QFontMetricsF(self.f_caption).horizontalAdvance("⚠ " + x) for x in self.w.problems),
        )
        if self.w.empty:
            if p:
                p.setFont(self.f_caption)
                p.drawText(QRectF(PAD, y, 600, 20), Qt.AlignmentFlag.AlignLeft, self.t.t("apw.empty"))
            return QSizeF(width, y + 40)
        for room in self.w.rooms:
            y, w = self._room(p, room, y)
            width = max(width, w)
            y += 26
        if self.w.problems:
            for text in self.w.problems:
                if p:
                    p.setPen(QColor("#C0392B"))
                    p.setFont(self.f_caption)
                    p.drawText(QRectF(PAD, y, 3000, 18), Qt.AlignmentFlag.AlignLeft, "⚠ " + text)
                y += 20
        return QSizeF(width + PAD, y + PAD)

    def _room(self, p: QPainter | None, room: WiringRoom, y: float) -> tuple[float, float]:
        s = self.s
        n_sw = len(room.switches)
        head = self.t.t("apw.room_head", room=room.title, aps=len(room.aps), sw=n_sw)
        right = PAD + QFontMetricsF(self.f_room).horizontalAdvance(head)
        if p:
            p.setPen(QColor(s.light_text))
            p.setFont(self.f_room)
            p.drawText(QRectF(PAD, y, 3000, 22), Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter, head)
        y += 32
        # the plan: APs row by row
        cols = room.columns
        top = y
        bottom = y
        for i, ap in enumerate(room.aps):
            row, col = (ap.row, ap.col) if cols > 0 else divmod(i, WRAP)
            x = PAD + col * (CELL_W + GAP)
            yy = top + row * (CELL_H + GAP)
            right = max(right, x + CELL_W)
            bottom = max(bottom, yy + CELL_H)
            if p:
                self._ap(p, QRectF(x, yy, CELL_W, CELL_H), ap.name, room, ap.switch, ap.port)
        y = bottom + 22
        # switch cards with their ports
        x = PAD
        card_bottom = y
        per_row = max(1, int((max(right, PAD + 3 * (CARD_W + 14)) - PAD + 14) // (CARD_W + 14)))
        for k, sw in enumerate(room.switches):
            if k and k % per_row == 0:
                x = PAD
                y = card_bottom + 14
            lines = max(1, (len(sw.used) + 1) // 2)
            h = 46 + lines * LINE_H + 8
            rect = QRectF(x, y, CARD_W, h)
            if p:
                self._card(p, rect, k, sw.label, f"{sw.model} · {sw.cabinet} · {len(sw.used)}/{sw.ports}", sw.used)
            right = max(right, x + CARD_W)
            card_bottom = max(card_bottom, y + h)
            x += CARD_W + 14
        return card_bottom, right

    def _ap(self, p: QPainter, r: QRectF, name: str, room: WiringRoom, sw: int, port: int) -> None:
        color = QColor(switch_color(sw))
        fill = QColor(color)
        fill.setAlphaF(0.16)
        p.setPen(QPen(color, 1.6))
        p.setBrush(fill)
        p.drawRoundedRect(r, 8, 8)
        p.setPen(QColor(self.s.light_text))
        p.setFont(self.f_name)
        fm = QFontMetricsF(self.f_name)
        p.drawText(
            r.adjusted(6, 5, -6, -r.height() / 2),
            Qt.AlignmentFlag.AlignCenter,
            fm.elidedText(name, Qt.TextElideMode.ElideLeft, r.width() - 12),
        )
        p.setPen(color.darker(115) if QColor(self.s.bg).lightnessF() > 0.5 else color.lighter(130))
        p.setFont(self.f_small)
        text = (
            f"{short_name(room.switches[sw].label)} · {self.t.t('apw.port', n=port)}"
            if sw >= 0
            else self.t.t("apw.none")
        )
        p.drawText(r.adjusted(4, r.height() / 2, -4, -4), Qt.AlignmentFlag.AlignCenter, text)

    def _card(self, p: QPainter, r: QRectF, k: int, title: str, sub: str, used: list[tuple[int, str]]) -> None:
        color = QColor(switch_color(k))
        p.setPen(QPen(color, 1.6))
        p.setBrush(QColor(self.s.bg))
        p.drawRoundedRect(r, 8, 8)
        head = QRectF(r.x(), r.y(), r.width(), 40)
        fill = QColor(color)
        fill.setAlphaF(0.18)
        p.setPen(Qt.PenStyle.NoPen)
        p.setBrush(fill)
        p.drawRoundedRect(head.adjusted(1, 1, -1, 0), 7, 7)
        p.setPen(QColor(self.s.light_text))
        p.setFont(self.f_name)
        p.drawText(head.adjusted(10, 3, -8, -18), Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter, title)
        p.setPen(QColor(self.s.muted))
        p.setFont(self.f_small)
        p.drawText(head.adjusted(10, 21, -8, -2), Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter, sub)
        p.setPen(QColor(self.s.light_text))
        half = (len(used) + 1) // 2
        for i, (port, ap) in enumerate(used):
            col, row = divmod(i, half) if half else (0, i)
            x = r.x() + 10 + col * (r.width() - 20) / 2
            y = r.y() + 46 + row * LINE_H
            p.drawText(
                QRectF(x, y, (r.width() - 20) / 2, LINE_H),
                Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter,
                f"{port:>2} → {short_name(ap)}",
            )


def export_ap_wiring_xlsx(wiring: ApWiring, path: str | Path, lang: str = "uk", title: str = "") -> Path:
    """The scheme as a table: one row per access point (room, AP, zone, switch, cabinet, port)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    t = Translator(lang)
    wb = Workbook()
    ws = wb.active
    ws.title = t.t("apw.sheet")[:31]
    row = 1
    if title:
        ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14)
        row = 3
    heads = [t.t(f"apw.col.{k}") for k in ("room", "ap", "zone", "switch", "cabinet", "port")]
    for c, h in enumerate(heads, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(vertical="center")
    for room in wiring.rooms:
        for ap in room.aps:
            row += 1
            sw = room.switches[ap.switch] if ap.switch >= 0 else None
            values = [room.title, ap.name, ap.zone, sw.label if sw else "—", sw.cabinet if sw else "", ap.port or None]
            for c, v in enumerate(values, 1):
                ws.cell(row=row, column=c, value=v)
            if sw is not None:
                color = switch_color(ap.switch).lstrip("#")
                ws.cell(row=row, column=4).font = Font(bold=True, color=color)
    for text in wiring.problems:
        row += 2
        ws.cell(row=row, column=1, value="⚠ " + text).font = Font(color="C0392B")
    for col, width in zip("ABCDEF", (34, 20, 26, 22, 10, 8), strict=True):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = ws.cell(row=(4 if title else 2), column=1)
    wb.save(str(path))
    return Path(path)
